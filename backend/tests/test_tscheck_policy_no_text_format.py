"""Long numeric Policy No is not mangled in exports: xlsx cell stays text (data_type 's',
number_format '@') with all digits intact; csv wraps long all-digit policy numbers as ="..." so
Excel keeps them as text."""
import csv
import io

from openpyxl import load_workbook

from .conftest import login, upload_and_wait

LONG_POLICY_NO = "500503312610001941"


def _patch_long_policy_no(c, pid: str) -> dict:
    doc = c.get(f"/policies/{pid}").json()
    fields = doc["fields"]
    fields["policy_number"] = LONG_POLICY_NO
    fields["gross_premium"] = fields.get("gross_premium") or 15642.08
    r = c.patch(f"/policies/{pid}", json={"fields": fields})
    assert r.status_code == 200, r.text
    return r.json()


def test_long_policy_number_preserved_as_text_in_xlsx_and_csv():
    c = login("agt1001", "Agent@123")
    doc = upload_and_wait(c)
    assert doc["status"] != "processing", f"OCR did not finish in time: {doc}"
    pid = doc["id"]
    patched = _patch_long_policy_no(c, pid)
    assert patched["fields"]["policy_number"] == LONG_POLICY_NO

    # --- xlsx ---
    resp = c.get("/mis/export", params={"format": "xlsx"})
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["Policy MIS"]
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    idx = next(i for i, cell in enumerate(header_row) if str(cell.value).strip() == "Policy No")
    target_row = next(
        r for r in ws.iter_rows(min_row=2) if r[idx].value is not None and str(r[idx].value) == LONG_POLICY_NO
    )
    cell = target_row[idx]
    assert cell.data_type == "s", cell.data_type
    assert cell.number_format == "@", cell.number_format
    assert str(cell.value) == LONG_POLICY_NO, cell.value

    # column must be wide enough to show the full number
    col_letter = header_row[idx].column_letter
    assert ws.column_dimensions[col_letter].width >= len(LONG_POLICY_NO), ws.column_dimensions[col_letter].width

    # --- csv ---
    resp_csv = c.get("/mis/export", params={"format": "csv"})
    assert resp_csv.status_code == 200, resp_csv.text
    raw = resp_csv.content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(raw)))
    header = [h.strip() for h in rows[0]]
    csv_idx = header.index("Policy No")
    target = next(row for row in rows[1:] if LONG_POLICY_NO in row[csv_idx])
    assert target[csv_idx] == f'="{LONG_POLICY_NO}"', target[csv_idx]
    c.close()
