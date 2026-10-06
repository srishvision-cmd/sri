"""MIS export criterion: CSV/XLSX export respects role scoping and filters.
agt1001 export contains only AGT1001 rows; admin export with agent filter restricts to that agent."""
import csv
import io

import httpx
from openpyxl import load_workbook

from .conftest import API_URL


def _login(username, password):
    c = httpx.Client(base_url=API_URL, timeout=30.0)
    r = c.post("/auth/login", json={"username": username, "password": password})
    assert r.status_code == 200, r.text
    return c


def test_agent_csv_export_scoped_to_own_policies():
    c = _login("agt1001", "Agent@123")
    resp = c.get("/mis/export", params={"format": "csv"})
    assert resp.status_code == 200, resp.text
    assert "csv" in resp.headers.get("content-type", "").lower() or resp.text.count(",") > 0
    rows = list(csv.reader(io.StringIO(resp.text)))
    assert len(rows) > 1  # header + at least one data row
    header = rows[0]
    agent_col = next((i for i, h in enumerate(header) if "agent" in h.lower()), None)
    assert agent_col is not None, header
    for row in rows[1:]:
        assert "AGT1001" in row[agent_col], row
    c.close()


def test_admin_xlsx_export_with_agent_filter_is_valid_xlsx():
    c = _login("admin", "Admin@123")
    agents_resp = c.get("/agents")
    assert agents_resp.status_code == 200, agents_resp.text
    agt1002 = next(a for a in agents_resp.json() if a["username"] == "agt1002")
    resp = c.get("/mis/export", params={"format": "xlsx", "agent_id": agt1002["id"]})
    assert resp.status_code == 200, resp.text
    content_type = resp.headers.get("content-type", "")
    assert "spreadsheet" in content_type or "octet-stream" in content_type, content_type
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) > 1
    header = [str(h) for h in rows[0]]
    agent_col = next((i for i, h in enumerate(header) if "agent" in h.lower()), None)
    assert agent_col is not None, header
    for row in rows[1:]:
        assert row[agent_col] and "AGT1002" in str(row[agent_col]), row
    c.close()
