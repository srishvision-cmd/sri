"""Criterion: a policy whose RTO office name (MARTHANDAM) is migrated to a TN code now matches
TN commission rules, both on GET /policies/{id}/commission (used by the detail panel) and in the
MIS Excel export's admin commission columns.

Read-only against the seeded fixture policy qa-rto-1 (policy no QA-RTO-1) + its dedicated seed
rules qa-rto-rule-receivable / qa-rto-rule-payable — never mutated here, only asserted on. The
server-startup migration (routers.policies.fix_rto_names, exercised in
test_tscheck_rto_normalisation.py) is what turns its stored rto from MARTHANDAM into TN75; this
file assumes that has already run (it runs unconditionally at boot).
"""
import io

import openpyxl

from ._commission_helpers import admin_client

POLICY_ID = "qa-rto-1"
POLICY_NO = "QA-RTO-1"


def test_commission_api_shows_values_for_migrated_rto_policy():
    admin = admin_client()
    r = admin.get(f"/policies/{POLICY_ID}")
    assert r.status_code == 200, r.text
    assert r.json()["fields"]["rto"] == "TN75", r.json()["fields"]

    r2 = admin.get(f"/policies/{POLICY_ID}/commission")
    assert r2.status_code == 200, r2.text
    body = r2.json()
    assert body["receivable"] == 2490, body
    assert body["payable"] == 2054.25, body
    assert body["margin"] == 435.75, body
    # no longer missing a TN-rto rule
    reasons = [reason for miss in (body.get("misses") or []) for reason in miss.get("reasons", [])]
    assert not any("RTO is MARTHANDAM" in reason for reason in reasons), reasons


def test_mis_export_xlsx_has_commission_values_for_migrated_rto_policy():
    admin = admin_client()
    r = admin.get("/mis/export", params={"format": "xlsx", "q": POLICY_NO})
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    rto_col = header.index("RTO")
    recv_col = header.index("Commission Receivable")
    pay_col = header.index("Commission Payable")
    margin_col = header.index("Commission Margin")
    policy_no_col = header.index("Policy No")

    row = next((row for row in ws.iter_rows(min_row=2, values_only=True) if str(row[policy_no_col]).strip("='\" ") == POLICY_NO), None)
    assert row is not None, "QA-RTO-1 row not found in export"
    assert row[rto_col] == "TN75", row
    assert row[recv_col] == 2490, row
    assert row[pay_col] == 2054.25, row
    assert row[margin_col] == 435.75, row
