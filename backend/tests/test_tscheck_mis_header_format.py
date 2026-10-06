"""MIS export criterion: the exported sheet's header row (A..AH + 2 payment columns) matches the
business reference file /tmp/mis_order.xlsx exactly (trimmed), for both xlsx and csv formats."""
import csv
import io

from openpyxl import load_workbook

from .conftest import login

REFERENCE_PATH = "/tmp/mis_order.xlsx"


def _reference_headers() -> list[str]:
    wb = load_workbook(REFERENCE_PATH)
    ws = wb.active
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return [str(c).strip() for c in row if c is not None]


def test_xlsx_export_sheet_and_header_match_reference():
    c = login("admin", "Admin@123")
    resp = c.get("/mis/export", params={"format": "xlsx"})
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Policy MIS" in wb.sheetnames, wb.sheetnames
    ws = wb["Policy MIS"]
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(c).strip() for c in row if c is not None]

    reference = _reference_headers()
    # exported sheet = reference business order (A..AH) followed by the two payment columns
    assert headers[: len(reference)] == reference, (headers[: len(reference)], reference)
    assert headers[len(reference):] == ["Amount Paid", "Agent Outstanding Premium"], headers[len(reference):]
    c.close()


def test_csv_export_header_matches_reference():
    c = login("admin", "Admin@123")
    resp = c.get("/mis/export", params={"format": "csv"})
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    header = next(csv.reader(io.StringIO(text)))
    header = [h.strip() for h in header]

    reference = _reference_headers()
    assert header[: len(reference)] == reference, (header[: len(reference)], reference)
    assert header[len(reference):] == ["Amount Paid", "Agent Outstanding Premium"], header[len(reference):]
    c.close()
