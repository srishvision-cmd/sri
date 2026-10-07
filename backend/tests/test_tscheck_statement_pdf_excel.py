"""Criterion: PDF & Excel.

GET /api/statements/{id}/pdf -> 200 application/pdf, text contains sections "A." and "B." and
"NET PAYABLE"; GET /api/statements/{id}/excel -> 200 xlsx, the commission-line policy number
cell is present as text in the sheet.
"""
import io

import fitz  # pymupdf
from openpyxl import load_workbook

from ._commission_helpers import admin_client
from ._statement_helpers import build_scenario, delete_statement, teardown_scenario


def test_pdf_has_sections_and_net_payable():
    admin = admin_client()
    scenario = build_scenario(admin)
    sid = None
    try:
        r = admin.post("/statements", json={"agent_id": scenario["agent"]["id"], "period_type": "daily"})
        assert r.status_code == 200, r.text
        s = r.json()
        sid = s["id"]
        policy_number = s["lines"][0]["policy_number"]

        pdf_r = admin.get(f"/statements/{sid}/pdf")
        assert pdf_r.status_code == 200, pdf_r.text
        assert pdf_r.headers["content-type"].startswith("application/pdf"), pdf_r.headers
        doc = fitz.open(stream=pdf_r.content, filetype="pdf")
        text = "\n".join(page.get_text() for page in doc)
        assert "A." in text, text
        assert "B." in text, text
        assert "NET PAYABLE" in text.upper(), text

        xl_r = admin.get(f"/statements/{sid}/excel")
        assert xl_r.status_code == 200, xl_r.text
        assert xl_r.headers["content-type"].startswith("application/vnd.openxmlformats"), xl_r.headers
        wb = load_workbook(io.BytesIO(xl_r.content))
        found = False
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None and str(cell) == policy_number:
                        found = True
        assert found, f"policy number {policy_number} not found as text in xlsx"
    finally:
        if sid:
            admin.delete(f"/statements/{sid}")
            delete_statement(sid)
        teardown_scenario(scenario)
