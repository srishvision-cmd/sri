"""Commission statement PDF + Excel rendering."""

import io
from lib.excel_safe import neutralise_formulas
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from models.statement import CommissionStatement

COMPANY = "Srish Vision Insurance"


def _rs(v: Optional[float]) -> str:
    return "-" if v is None else f"Rs {v:,.2f}"


def _period(s: CommissionStatement) -> str:
    return f"{s.period_type.title()} · {s.from_date} to {s.to_date}"


def _split(s: CommissionStatement) -> tuple[float, float]:
    """(cash pending, partial payment pending) out of the outstanding lines."""
    pp = round(sum(o.outstanding for o in s.outstanding_lines if o.payment_mode == "Partial Payment"), 2)
    return round(s.total_outstanding - pp, 2), pp


def _net_text(s: CommissionStatement) -> tuple[str, float]:
    net = round(s.total_payable - s.total_outstanding, 2)
    return ("Payable commission — Srishvision pays partner" if net >= 0 else "Receivable — partner pays Srishvision"), abs(net)


def statement_pdf(s: CommissionStatement) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Commission statement {s.number}")
    st = getSampleStyleSheet()
    small = st["BodyText"].clone("small", fontSize=8, leading=10)
    head = colors.HexColor("#1E3A5F")

    def table(rows, widths, money_cols=()):
        t = Table(rows, colWidths=widths, repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), head), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0BEC5")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F3F6F9")]),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ]
        style += [("ALIGN", (c, 0), (c, -1), "RIGHT") for c in money_cols]
        t.setStyle(TableStyle(style))
        return t

    story = [
        Paragraph(f"<b>{COMPANY}</b> — Partner Payout Statement", st["Title"]),
        Paragraph(
            f"<b>Statement:</b> {s.number} &nbsp;&nbsp; <b>Partner:</b> {s.agent_name} ({s.agent_code}) &nbsp;&nbsp; <b>Branch:</b> {s.branch or '-'} "
            f"&nbsp;&nbsp; <b>Period:</b> {_period(s)} &nbsp;&nbsp; <b>Generated:</b> {s.created_at:%d-%m-%Y} by {s.created_by_name}",
            st["BodyText"],
        ),
        Spacer(1, 4 * mm),
        Paragraph("<b>A. Commission earned</b>", st["Heading3"]),
    ]
    rows = [["#", "Policy No", "Reg No", "Insured", "Insurer", "Product", "Risk start", "Net premium", "Commission %", "Commission"]]
    for i, l in enumerate(s.lines, 1):
        rows.append([i, Paragraph(l.policy_number or "-", small), l.reg_no or "-", Paragraph(l.insured_name or "-", small), Paragraph(l.insurer or "-", small),
                     Paragraph(l.product or "-", small), l.start_date or "-", _rs(l.net_premium), Paragraph(l.commission_rate or "-", small), _rs(l.payable)])
    rows.append(["", "", "", "", "", "", "", "", "Total", _rs(s.total_payable)])
    story.append(table(rows, [8 * mm, 40 * mm, 24 * mm, 36 * mm, 32 * mm, 30 * mm, 19 * mm, 26 * mm, 31 * mm, 26 * mm], (7, 9)))

    story += [Spacer(1, 4 * mm), Paragraph("<b>B. Premium outstanding deducted</b>", st["Heading3"])]
    rows = [["#", "Policy No", "Reg No", "Insured", "Mode", "Total premium", "Outstanding", "Deducted", "Still outstanding"]]
    for i, o in enumerate(s.outstanding_lines, 1):
        rows.append([i, Paragraph(o.policy_number or "-", small), o.reg_no or "-", Paragraph(o.insured_name or "-", small), o.payment_mode or "-",
                     _rs(o.gross_premium), _rs(o.outstanding), _rs(o.deducted), _rs(round(o.outstanding - o.deducted, 2))])
    rows.append(["", "", "", "", "", "Total", _rs(s.total_outstanding), _rs(s.outstanding_deducted), _rs(s.outstanding_balance)])
    t_out = table(rows, [8 * mm, 44 * mm, 26 * mm, 48 * mm, 24 * mm, 28 * mm, 28 * mm, 28 * mm, 30 * mm], (5, 6, 7, 8))
    for i, o in enumerate(s.outstanding_lines, 1):  # not (fully) recovered: commission was not enough
        if o.outstanding - o.deducted > 0.005:
            t_out.setStyle(TableStyle([("TEXTCOLOR", (8, i), (8, i), colors.HexColor("#B91C1C"))]))
    for i, o in enumerate(s.outstanding_lines, 1):  # highlight Partial Payment cases
        if o.payment_mode == "Partial Payment":
            t_out.setStyle(TableStyle([("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFF3C4"))]))
    story.append(t_out)
    cash_p, part_p = _split(s)
    net_label, net_abs = _net_text(s)

    summary = [
        ["Summary", ""],
        ["Total commission (A)", _rs(s.total_payable)],
        ["Less: total premium pending (Cash cases)", _rs(cash_p)],
        ["Less: partial payment pending", _rs(part_p)],
        [net_label, _rs(net_abs)],
        ["Premium outstanding deducted now (B)", _rs(s.outstanding_deducted)],
        ["Outstanding balance carried forward", _rs(s.outstanding_balance)],
        ["NET PAYABLE TO AGENT", _rs(s.net_payable)],
    ]
    if s.status == "paid":
        summary.append(["Paid on / UTR", f"{s.paid_date} / {s.utr}"])
    t = Table(summary, colWidths=[80 * mm, 50 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), head), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 4), (-1, 4), "Helvetica-Bold"), ("BACKGROUND", (0, 3), (-1, 3), colors.HexColor("#FFF3C4")),
        ("FONTNAME", (0, 7), (-1, 7), "Helvetica-Bold"), ("BACKGROUND", (0, 7), (-1, 7), colors.HexColor("#E0F2F1")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B0BEC5")), ("ALIGN", (1, 1), (1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 9),
    ]))
    story += [Spacer(1, 5 * mm), t, Spacer(1, 4 * mm),
              Paragraph(f"Status: <b>{'PAID' if s.status == 'paid' else 'PENDING PAYMENT'}</b>. Computer generated statement.", small)]
    doc.build(story)
    return buf.getvalue()


def statement_xlsx(s: CommissionStatement) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="1E3A5F")
    white = Font(bold=True, color="FFFFFF")

    def header(row: list):
        ws.append(row)
        for c in ws[ws.max_row]:
            c.font, c.fill = white, fill

    ws.append([f"{COMPANY} — Partner Payout Statement"])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in (("Statement", s.number), ("Partner", f"{s.agent_name} ({s.agent_code})"), ("Branch", s.branch or "-"), ("Period", _period(s)),
                 ("Status", "PAID" if s.status == "paid" else "PENDING PAYMENT"), ("Paid on", s.paid_date or "-"), ("UTR", s.utr or "-")):
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    ws.append(["A. Commission earned"])
    ws.cell(ws.max_row, 1).font = bold
    header(["#", "Policy No", "Reg No", "Insured", "Insurer", "Product", "Risk start", "Net premium", "Commission %", "Commission", "Rule"])
    for i, l in enumerate(s.lines, 1):
        ws.append([i, l.policy_number, l.reg_no, l.insured_name, l.insurer, l.product, l.start_date, l.net_premium, l.commission_rate or "-", l.payable, l.rule_name])
    ws.append(["", "", "", "", "", "", "", "", "Total", s.total_payable])
    ws.cell(ws.max_row, 10).font = bold
    ws.append([])
    ws.append(["B. Premium outstanding deducted"])
    ws.cell(ws.max_row, 1).font = bold
    header(["#", "Policy No", "Reg No", "Insured", "Mode", "Total premium", "Outstanding", "Deducted", "Still outstanding"])
    amber = PatternFill("solid", fgColor="FFF3C4")
    for i, o in enumerate(s.outstanding_lines, 1):
        ws.append([i, o.policy_number, o.reg_no, o.insured_name, o.payment_mode, o.gross_premium, o.outstanding, o.deducted, round(o.outstanding - o.deducted, 2)])
        if o.payment_mode == "Partial Payment":
            for c in ws[ws.max_row]:
                c.fill = amber
    ws.append(["", "", "", "", "", "Total", s.total_outstanding, s.outstanding_deducted])
    ws.append([])
    cash_p, part_p = _split(s)
    net_label, net_abs = _net_text(s)
    ws.append(["SUMMARY"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    for k, v in (("Total commission (A)", s.total_payable), ("Less: total premium pending (Cash cases)", cash_p), ("Less: partial payment pending", part_p),
                 (net_label, net_abs), ("Premium outstanding deducted now (B)", s.outstanding_deducted),
                 ("Outstanding balance carried forward", s.outstanding_balance), ("NET PAYABLE TO AGENT", s.net_payable)):
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
        if k.startswith("Less: partial"):
            ws.cell(ws.max_row, 1).fill = amber
    for col, w in zip("ABCDEFGHIJK", (34, 26, 14, 28, 26, 24, 14, 16, 22, 16, 28)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
            elif c.column == 2 and c.value is not None:
                c.number_format = "@"  # long policy numbers stay exact text
    out = io.BytesIO()
    neutralise_formulas(wb)
    wb.save(out)
    return out.getvalue()
