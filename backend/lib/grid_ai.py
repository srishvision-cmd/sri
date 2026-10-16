"""Read a commission grid (Excel / PDF / image) with Gemini into rows in the grid-Excel column format."""
import asyncio
import io
import json
import os
import re

from emergentintegrations.llm.chat import FileContentWithMimeType, LlmChat, UserMessage
from openpyxl import load_workbook

from models.commission import AiDeclinedRow, AiGridRow

MODEL_PROVIDER, MODEL_NAME = "gemini", "gemini-3-flash-preview"
TIMEOUT_S = 150
MAX_SHEET_CELLS = 40000

SYSTEM = ("You convert Indian general-insurance commission grids (payout / rate cards from insurers or channel partners) "
          "into structured rules. You only output strict JSON.")


def _prompt(kind: str, hints: dict) -> str:
    """Kept short on purpose — long instructions make the model think for minutes."""
    return f"""Extract every commission rate cell from this grid as JSON only:
{{"note": "company / channel partner / period you saw", "rules": [{{"kind": "{kind.title() if kind else 'Receivable'}", "insurer": "", "broker": "", "ptype": "", "products": "",
"rtos": "", "vehicle_ages": "", "fuel_types": "", "business_types": "", "ncb": "", "cover": "", "gvw_min": 0, "gvw_max": 0, "cc_min": 0, "cc_max": 0, "seating_min": 0, "seating_max": 0,
"od_percent": 0, "tp_percent": 0, "net_percent": 0, "odtp_percent": 0, "flat_amount": 0, "valid_from": "YYYY-MM-DD", "valid_to": "YYYY-MM-DD", "needs_review": false, "review_note": ""}}]}}
{f'Channel partner for every rule: {hints["broker"]}.' if hints.get("broker") else ""}
Rules: one row per (slab × age × plan) combination; omit keys that are not stated; numbers without % or commas; blank or "-" = no rule; never guess.
"on OD" → od_percent, "on TP" → tp_percent, "on Net" → net_percent, "OD+TP"/"on total" → odtp_percent, rupees → flat_amount. "above 12000 kg" → gvw_min 12001.
vehicle_ages must use these labels, comma separated: {", ".join(hints["ages"])}. List EVERY label a range covers:
"New" → New Vehicle; "Other than New upto 10 years" / "1-10 yrs" → 1-5 Years, 6-7 Years, 8-10 Years; ">10 to <=15" → 11-12 Years, 13-15 Years; ">15" → Above 15 Years.
Manufacturing years (e.g. "2019 & older") → convert to age bands as of the grid period.
GVW / CC / seating limits: "<=2000" or "upto 2T" → gvw_max 2000; "> 2000 <=3500" → gvw_min 2001, gvw_max 3500; ">40000" → gvw_min 40001 (tonnes × 1000). Fill them on EVERY row of that slab.
When the grid splits rates into parts (e.g. Remuneration + Reward) use the TOTAL commission column. ncb: "with" / "without" only when the grid gives separate rates for with-NCB and without-NCB (zero NCB) policies, else omit.
cover (PCV, GCV and Misc D only): "Package" for package / comprehensive / OD rates, "SATP" for stand-alone TP / liability-only / act-only rates, blank when the rate applies to both. When the grid has separate Package and SATP columns or sections, output one rule each.
ptype one of: {", ".join(hints["ptypes"])}.
products (comma separated) from: {", ".join(hints["products"])}. insurer: use the closest of {", ".join(hints["insurers"][:40])}.
valid_from / valid_to: only when the file prints explicit start / end dates ("valid from 01-08-2026 to 31-08-2026"); a month or period title like "Aug 2026" alone is NOT a date range — omit both.
Set needs_review true (with review_note) when a value is blurry, merged across cells or ambiguous.
Also return "declined": [{{"ptype": "", "products": "", "rtos": "", "makes": "", "models": "", "cc_min": 0, "cc_max": 0, "gvw_min": 0, "gvw_max": 0, "note": ""}}] — one item per row of any
"Declined" / "Excluded" / "Not allowed" / "Negative list" table or remark (RTOs, CC, GVW, makes or models the company does not pay for). rtos as VAHAN codes
(e.g. TN-29 → TN29), comma separated; makes / models as printed; "above 1500 cc" → cc_min 1501; GVW in kg ("above 40 T" → gvw_min 40001, "upto 2.5 T" → gvw_max 2500). Group rows with the same product + remark into one item.
Use the RTO Name column only to fill a missing code. Never put declined rows in "rules". Omit "declined" when there is no such table.
{hints.get("template") or ""}"""


def _xlsx_text(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out, cells = [], 0
    for ws in wb.worksheets:
        out.append(f"### Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in row]
            if any(vals):
                out.append(" | ".join(vals).rstrip(" |"))
                cells += len(vals)
            if cells > MAX_SHEET_CELLS:
                out.append("(truncated)")
                return "\n".join(out)
    return "\n".join(out)


def _parse(text: str) -> dict:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("The AI could not find a commission grid in this file")
    return json.loads(text[start:end + 1])


async def read_grid(path: str, mime: str, file_name: str, kind: str, hints: dict, session_id: str) -> tuple[list[AiGridRow], str, list[AiDeclinedRow]]:
    return await asyncio.wait_for(asyncio.to_thread(_run_in_own_loop, path, mime, file_name, kind, hints, session_id), timeout=TIMEOUT_S + 15)


def _run_in_own_loop(*args) -> tuple[list[AiGridRow], str, list[AiDeclinedRow]]:
    """Worker-thread event loop (the LLM SDK blocks). Not asyncio.run(): its executor shutdown can wait forever on SDK threads."""
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(_read(*args))
    finally:
        loop.close()


async def _read(path: str, mime: str, file_name: str, kind: str, hints: dict, session_id: str) -> tuple[list[AiGridRow], str, list[AiDeclinedRow]]:
    chat = LlmChat(api_key=os.environ["EMERGENT_LLM_KEY"], session_id=session_id, system_message=SYSTEM).with_model(MODEL_PROVIDER, MODEL_NAME)
    prompt = _prompt(kind, hints)
    if file_name.lower().endswith((".xlsx", ".xlsm")):
        with open(path, "rb") as fh:
            msg = UserMessage(text=f"{prompt}\n\nGRID (Excel, cells separated by |):\n{_xlsx_text(fh.read())}")
    else:
        msg = UserMessage(text=prompt, file_contents=[FileContentWithMimeType(mime_type=mime, file_path=path)])
    try:
        text = await asyncio.wait_for(chat.send_message(msg), timeout=TIMEOUT_S)
    except asyncio.TimeoutError:
        raise RuntimeError("Reading the grid timed out — try a smaller file or split the pages")
    data = _parse(text)
    rows = []
    for r in data.get("rules") or []:
        if not isinstance(r, dict):
            continue
        clean = {k: (", ".join(map(str, v)) if isinstance(v, list) else v) for k, v in r.items() if k in AiGridRow.model_fields}
        try:
            row = AiGridRow(**clean)
        except ValueError:
            row = AiGridRow(**{k: v for k, v in clean.items() if isinstance(v, str)}, needs_review=True, review_note="Some numbers could not be read")
        if not any(getattr(row, k) is not None for k in ("od_percent", "tp_percent", "net_percent", "odtp_percent", "flat_amount")):
            row.needs_review, row.review_note = True, row.review_note or "No rate found in this row"
        rows.append(row)
    declined = []
    for d in data.get("declined") or []:
        if not isinstance(d, dict):
            continue
        clean = {k: (", ".join(map(str, v)) if isinstance(v, list) else v) for k, v in d.items() if k in AiDeclinedRow.model_fields}
        clean = {k: v for k, v in clean.items() if v not in (None, "", 0)}
        for k in ("gvw_min", "gvw_max"):  # tonnes slipped through (e.g. 40 → 40000 kg)
            if isinstance(clean.get(k), (int, float)) and clean[k] < 100:
                clean[k] = clean[k] * 1000
        try:
            item = AiDeclinedRow(**clean)
        except ValueError:
            item = AiDeclinedRow(**{k: v for k, v in clean.items() if isinstance(v, str)})
        if item.rtos or item.makes or item.models or item.cc_min or item.cc_max or item.gvw_min or item.gvw_max:
            declined.append(item)
    return rows, str(data.get("note") or "")[:300], declined
