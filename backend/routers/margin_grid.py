"""Group margin grid: company × product rows, one column per agent group, cell = points below the receivable (−2, −3…).

Payable for agents in the group = receivable % − margin, so it follows every receivable rule change automatically.
New margins apply from a date (old policies keep the old margin) unless "apply to past policies" is ticked.
"""
import io
from datetime import date, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from lib.auth import require_roles
from lib.commission import chain_margins, product_type_of, single_product
from lib.db import db
from models.commission import MarginGrid, MarginGridIn, MarginGridRow, MarginGridSaveResult
from models.masters import PRODUCT_TYPES
from models.user import User
from routers.commission import _today, commission_masters

router = APIRouter(tags=["margin-grid"])
admin_only = require_roles("admin")
ALL_COS, ALL_PRODUCTS = "All other companies", "All products"


def _key(m: dict) -> tuple:
    return (m.get("insurer") or None, m.get("product_type") or None, m.get("product") or None)


def _latest(lines: list[dict]) -> Optional[dict]:
    """The line that applies going forward (latest start); None when that line was stopped."""
    if not lines:
        return None
    last = max(lines, key=lambda m: m.get("valid_from") or "")
    return None if last.get("ended") else last


async def _groups() -> list[dict]:
    return await db.agent_groups.find({}, {"_id": 0}).sort("name", 1).to_list(200)


async def _grid() -> MarginGrid:
    groups = await _groups()
    rows: dict[tuple, MarginGridRow] = {}
    for g in groups:
        by_key: dict[tuple, list] = {}
        for m in g.get("margins") or []:
            by_key.setdefault(_key(m), []).append(m)
        for k, lines in by_key.items():
            row = rows.setdefault(k, MarginGridRow(insurer=k[0], product_type=k[1], product=k[2]))
            cur = _latest(lines)
            row.margins[g["name"]] = cur["margin"] if cur else None
            row.since[g["name"]] = cur.get("valid_from") if cur else None
    order = sorted(rows.values(), key=lambda r: (r.insurer is None, (r.insurer or "").lower(), r.product_type is None, r.product_type or "", r.product or ""))
    return MarginGrid(groups=[g["name"] for g in groups], rows=order)


@router.get("/commission/margin-grid", response_model=MarginGrid)
async def margin_grid(_: User = Depends(admin_only)):
    return await _grid()


@router.get("/commission/margin-grid/suggest", response_model=list[MarginGridRow])
async def margin_grid_suggest(_: User = Depends(admin_only)):
    """Company × product type lines found in the live receivable grid — pre-fills the margin grid so nothing is typed."""
    today = _today()
    recs = await db.commission_rules.find({"kind": "receivable", "active": True, "superseded_by": None, "$or": [{"valid_to": None}, {"valid_to": {"$gte": today}}]},
                                          {"_id": 0, "insurer": 1, "products": 1}).to_list(5000)
    seen = {(r.get("insurer"), product_type_of(r.get("products") or []), single_product(r.get("products") or []) if product_type_of(r.get("products") or []) == "PCV" else None) for r in recs}
    return [MarginGridRow(insurer=i, product_type=t, product=p) for i, t, p in sorted(seen, key=lambda x: (x[0] is None, x[0] or "", x[1] or "", x[2] or ""))]


def _clean_row(r: MarginGridRow) -> MarginGridRow:
    r.insurer = (r.insurer or "").strip() or None
    r.product_type = (r.product_type or "").strip() or None
    r.product = (r.product or "").strip() or None
    if r.product and not r.product_type:
        r.product_type = next((t for t, items in PRODUCT_TYPES.items() if r.product in items), None)
    if r.product_type and r.product_type not in PRODUCT_TYPES:
        raise HTTPException(status_code=400, detail=f"Unknown product type '{r.product_type}'")
    if r.product and r.product not in PRODUCT_TYPES.get(r.product_type or "", []):
        raise HTTPException(status_code=400, detail=f"'{r.product}' is not a {r.product_type} product")
    for g, v in r.margins.items():
        if v is not None:
            v = abs(v)  # −2 and 2 both mean 2 points below the receivable
            if v > 100:
                raise HTTPException(status_code=400, detail=f"Margin {v:g} for Group {g} must be 0–100")
            r.margins[g] = v
    return r


@router.put("/commission/margin-grid", response_model=MarginGridSaveResult)
async def save_margin_grid(body: MarginGridIn, _: User = Depends(admin_only)):
    eff = (body.valid_from or "")[:10] or _today()
    try:
        date.fromisoformat(eff)
    except ValueError:
        raise HTTPException(status_code=400, detail="Applicable from must be a date (YYYY-MM-DD)")
    rows = [_clean_row(r) for r in body.rows]
    keys = [(r.insurer, r.product_type, r.product) for r in rows]
    if len(set(keys)) != len(keys):
        dup = next(k for k in keys if keys.count(k) > 1)
        raise HTTPException(status_code=400, detail=f"{dup[0] or ALL_COS} · {dup[2] or dup[1] or ALL_PRODUCTS} is listed twice")
    changed = 0
    stop = (date.fromisoformat(eff) - timedelta(days=1)).isoformat()
    for g in await _groups():
        lines = [dict(m) for m in g.get("margins") or []]
        touched = False
        wanted = {k: r.margins.get(g["name"]) for k, r in zip(keys, rows)}
        # rows removed from the grid = margin cleared for every group
        for k in {_key(m) for m in lines} - set(wanted):
            wanted[k] = None
        for k, v in wanted.items():
            mine = [m for m in lines if _key(m) == k]
            cur = _latest(mine)
            if (cur["margin"] if cur else None) == v:
                continue
            touched, changed = True, changed + 1
            others = [m for m in lines if _key(m) != k]
            if body.apply_to_past:  # one line for all policies; history replaced
                mine = [] if v is None else [{"insurer": k[0], "product_type": k[1], "product": k[2], "margin": v, "valid_from": None, "valid_to": None, "ended": False}]
            else:  # old policies keep their margin: replace lines starting on/after the date, stop or follow on with the new value
                mine = [m for m in mine if not m.get("valid_from") or m["valid_from"] < eff]
                if v is None:
                    if mine:  # every remaining line starts before the date → stop the latest the day before
                        max(mine, key=lambda m: m.get("valid_from") or "").update(ended=True, valid_to=stop)
                else:
                    mine.append({"insurer": k[0], "product_type": k[1], "product": k[2], "margin": v, "valid_from": eff, "valid_to": None, "ended": False})
            lines = others + mine
        if touched:
            await db.agent_groups.update_one({"id": g["id"]}, {"$set": {"margins": chain_margins(lines)}})
    return MarginGridSaveResult(changed=changed)


# ---------- Excel: same grid, only the numbers are typed ----------
@router.get("/commission/margin-grid/excel")
async def margin_grid_excel(_: User = Depends(admin_only)):
    grid = await _grid()
    suggest = await margin_grid_suggest(None)  # type: ignore[arg-type]
    have = {(r.insurer, r.product_type, r.product) for r in grid.rows}
    rows = grid.rows + [r for r in suggest if (r.insurer, r.product_type, r.product) not in have]
    wb = Workbook()
    ws = wb.active
    ws.title = "Margins"
    ws.append(["Insurance Company", "Product Type", "Product (optional)", *[f"Group {g}" for g in grid.groups]])
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F766E"), Alignment(horizontal="center")
    for r in rows:
        ws.append([r.insurer or ALL_COS, r.product_type or ALL_PRODUCTS, r.product or "", *[(-r.margins[g] if r.margins.get(g) is not None else None) for g in grid.groups]])
    ws.freeze_panes = "D2"
    for col, w in zip("ABC", (34, 16, 30)):
        ws.column_dimensions[col].width = w
    lists = wb.create_sheet("Lists")
    m = await commission_masters(None)  # type: ignore[arg-type]
    cols = {"Companies": [ALL_COS, *m.insurers], "Product Types": [ALL_PRODUCTS, *PRODUCT_TYPES], "Products": [p for v in PRODUCT_TYPES.values() for p in v]}
    for ci, (title, vals) in enumerate(cols.items(), start=1):
        lists.cell(1, ci, title).font = Font(bold=True)
        for ri, v in enumerate(vals, start=2):
            lists.cell(ri, ci, v)
        col = lists.cell(1, ci).column_letter
        dv = DataValidation(type="list", formula1=f"Lists!${col}$2:${col}${len(vals) + 1}", allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{'ABC'[ci - 1]}2:{'ABC'[ci - 1]}1000")
    info = wb.create_sheet("Instructions")
    for line in ("Type the margin per group as -2, -3, -4 (points below the receivable). Blank = no auto payable for that group.",
                 "Company, Product Type and Product are dropdowns. Product is optional — use it only for a specific product like 'PCV - Taxi'.",
                 "Upload this file on the Payable margins screen: it fills the grid there, you check it and press Save."):
        info.append([line])
    info.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="payable_margin_grid_{_today()}.xlsx"'})


@router.post("/commission/margin-grid/excel", response_model=MarginGrid)
async def margin_grid_read_excel(file: UploadFile = File(...), _: User = Depends(admin_only)):
    """Read the Excel back into grid rows (nothing saved — the screen shows them for Save)."""
    try:
        ws = load_workbook(io.BytesIO(await file.read()), data_only=True, read_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file")
    if not rows or str(rows[0][0] or "").strip().lower() != "insurance company":
        raise HTTPException(status_code=400, detail="Use the margin grid Excel downloaded from this screen")
    groups = [g["name"] for g in await _groups()]
    cols = {i: str(h).replace("Group", "").strip() for i, h in enumerate(rows[0]) if i >= 3 and h}
    unknown = [g for g in cols.values() if g not in groups]
    if unknown:
        raise HTTPException(status_code=400, detail=f"Unknown group column(s): {', '.join(unknown)}")
    out: list[MarginGridRow] = []
    for n, row in enumerate(rows[1:], start=2):
        if not row or all(v in (None, "") for v in row):
            continue
        cell = lambda i, row=row: str(row[i]).strip() if i < len(row) and row[i] not in (None, "") else ""  # noqa: E731
        margins = {}
        for i, g in cols.items():
            v = row[i] if i < len(row) else None
            if v in (None, ""):
                margins[g] = None
                continue
            try:
                margins[g] = abs(float(str(v).replace("%", "").strip()))
            except ValueError:
                raise HTTPException(status_code=400, detail=f"Row {n}: '{v}' for Group {g} is not a number")
        r = MarginGridRow(insurer=None if cell(0) in ("", ALL_COS) else cell(0), product_type=None if cell(1) in ("", ALL_PRODUCTS) else cell(1),
                          product=cell(2) or None, margins=margins)
        try:
            out.append(_clean_row(r))
        except HTTPException as e:
            raise HTTPException(status_code=400, detail=f"Row {n}: {e.detail}")
    return MarginGrid(groups=groups, rows=out)
