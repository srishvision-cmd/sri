"""Commission grid Excel template: download (blank or pre-filled with the current grid) and bulk upload.

One row = one rule. A row whose Rule ID (or, when blank, exact conditions) matches a current rule
updates that rule's rate from the row's Applicable From date (versioned); other rows add new rules.
"""
import io
from lib.excel_safe import neutralise_formulas
from lib.grid_templates import find_template, learn, prompt_block
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from lib.auth import require_roles
from lib.db import db
from typing import List

from lib.commission import LABEL, PCT_KEYS, group_margin_line, margin_text, single_product, check_against, derive_payable, overlaps, product_type_of
from models.commission import (
    AiGridApplyIn, AiGridRead,
    COMPONENTS, BulkPayableIn, BulkPayableResult, BulkPayableRow, CommissionRule, CommissionRuleIn, CommissionSettings, CompareRates, ComparePayable, CompareRow, GridRowError, GridRowPreview, GridUploadResult,
)
from models.masters import FUEL_BY_TYPE, FUEL_TYPES, MAKES, PRODUCT_TYPES, RTO_STATES, RTO_ZONES
from models.policy import AGE_BANDS, BUSINESS_TYPES
from models.user import User
from routers.commission import _clean, _signature, _today, apply_rule_update, commission_masters, insert_rule

router = APIRouter(tags=["commission-grid"])
admin_only = require_roles("admin")

# (header, key) — headers are what the admin sees; keys map to rule fields
COLS: list[tuple[str, str]] = [
    ("Rule ID (blank = new rule)", "id"), ("Type", "kind"), ("Rule Name (blank = auto)", "name"), ("Insurance Company", "insurer"),
    ("Channel Partners (comma separated, blank = all)", "broker"), ("Agent Group (payable)", "agent_group"), ("Agent (payable) — Code · Name, pick from list", "agent_code"),
    ("Product Type", "ptype"), ("Products (comma separated, blank = all of type)", "products"), ("Makes (comma separated)", "makes"),
    ("Model", "model"), ("RTOs (comma separated e.g. TN19,TN20 or TN)", "rtos"), ("Vehicle Age Bands (comma separated)", "vehicle_ages"),
    ("Fuel Types (comma separated)", "fuel_types"), ("Business Type (New, Renewal, Rollover, Used — blank = all)", "business_types"), ("NCB", "ncb"), ("Package / SATP", "cover"),
    ("GVW From (kg)", "gvw_min"), ("GVW To (kg)", "gvw_max"), ("CC From", "cc_min"), ("CC To", "cc_max"),
    ("Seating From", "seating_min"), ("Seating To", "seating_max"),
    ("OD %", "od_percent"), ("TP %", "tp_percent"), ("Net %", "net_percent"), ("OD+TP %", "odtp_percent"), ("Flat Amount", "flat_amount"),
    ("Applicable From (YYYY-MM-DD)", "valid_from"), ("Applicable To (YYYY-MM-DD)", "valid_to"), ("Active", "active"),
]
NUMS = ("gvw_min", "gvw_max", "cc_min", "cc_max", "seating_min", "seating_max", *COMPONENTS)
NCB_OUT = {"any": "Any", "with": "With NCB", "without": "Without NCB"}
COVER_OUT = {"any": "Any", "package": "Package", "satp": "SATP"}
ALL_PRODUCTS = [p for items in PRODUCT_TYPES.values() for p in items]
# product sheets: which product-specific columns each sheet shows (all other columns are shared)
_DIMS = {"gvw_min", "gvw_max", "cc_min", "cc_max", "seating_min", "seating_max"}
_SHEET_EXTRA = {"Private Car": {"cc_min", "cc_max"}, "Two Wheeler": {"cc_min", "cc_max"}, "PCV": {"seating_min", "seating_max", "cover"},
                "GCV": {"gvw_min", "gvw_max", "cover"}, "Misc D": {"cover"}}
_NON_MOTOR_DROP = {"makes", "model", "rtos", "vehicle_ages", "fuel_types", "ncb", "cover", *_DIMS}
ALL_SHEET = "All products"  # rules covering several product types (full column set incl. Product Type)
SKIP_SHEETS = {"Lists", "Instructions"}
MULTI = {"broker", "products", "makes", "rtos", "vehicle_ages", "fuel_types", "business_types"}  # comma-separated, dropdown is a helper


def sheet_cols(sheet: str) -> list[tuple[str, str]]:
    if sheet not in PRODUCT_TYPES:
        return COLS
    motor = sheet in _SHEET_EXTRA
    keep = lambda k: k != "ptype" and (k not in _DIMS and k != "cover" or k in _SHEET_EXTRA.get(sheet, set())) and (motor or k not in _NON_MOTOR_DROP)  # noqa: E731
    return [(h, k) for h, k in COLS if keep(k)]


def zone_label(z: dict) -> str:
    """Saved RTO group as shown in the Excel dropdown: 'Chennai' or 'Chennai (Shriram General)'."""
    return f"{z['name']} ({z['insurer']})" if z.get("insurer") else z["name"]


def rto_options(groups: list[dict]) -> list[str]:
    return [zone_label(z) for z in groups] + [n for n, _ in RTO_ZONES] + [f"{c} · {n}" for c, n, _ in RTO_STATES]


def _rtos(v, groups: list[dict]) -> list[str]:
    """'Chennai RTOs, My group, TN19, TN · Tamil Nadu' → plain RTO codes (group / zone names expand to their RTOs)."""
    zones = {n.lower(): codes for n, codes in RTO_ZONES} | {zone_label(z).lower(): z["rtos"] for z in groups} \
        | {z["name"].lower(): z["rtos"] for z in groups if not z.get("insurer")} | {n.lower(): [c] for c, n, _ in RTO_STATES}
    out: list[str] = []
    for tok in _split(v):
        out += zones.get(tok.lower()) or [tok.split("·")[0].strip().upper().replace(" ", "")]
    return list(dict.fromkeys(out))


INSTRUCTIONS = [
    "HOW TO USE THIS TEMPLATE",
    "1. One sheet per product (Private Car, Two Wheeler, PCV, GCV, Misc D, Health). One row = one commission rule.",
    "   Type = Receivable (from channel partner) or Payable (to agent). 'All products' sheet = rules for several product types.",
    "   Every criteria column has a dropdown. Columns that take several values (Channel Partners, Products, Makes, RTOs, Age, Fuel, Business Type):",
    "   pick one from the dropdown, or type several separated by commas (e.g. Fazil, Gourav  /  Chennai RTOs, TN19).",
    "   RTOs: pick a saved RTO group (Commission Grid → RTO groups), a zone (Chennai RTOs, ROTN), a whole state (TN · Tamil Nadu) or type codes (TN19, TN20).",
    "2. Leave a condition blank to apply to all (all companies, all channel partners, all RTOs, all makes...).",
    "3. Payables: copy a row and change only the Agent / Agent Group and rates to give different agents a different grid for the same company + channel partner.",
    "   Agent column shows 'CODE · Name' (e.g. AGT1022 · Sambu) — pick from the dropdown. Agent Group takes a group name only (no margin numbers).",
    "   Fill either Agent or Agent Group (or neither = common payable for all agents).",
    "4. To change an existing rate: download 'Current grid', edit the % and set Applicable From to the date the new rate starts. Policies before that date keep the old rate.",
    "   If you change only the rate and keep the old date, the new rate applies from today.",
    "5. New rules: leave Rule ID blank. A row with exactly the same conditions as an existing rule updates that rule instead of adding a duplicate.",
    "6. PCV rules need Products filled in. GCV rules need GVW From and/or GVW To.",
    "7. Lists of valid values are in the 'Lists' sheet.",
    "8. Receivable rows are processed before payable rows; a payable can never exceed the matching receivable.",
]


def _split(v) -> list[str]:
    return [x.strip() for x in str(v or "").replace(";", ",").split(",") if x.strip()]


def _num(v) -> Optional[float]:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace("%", "").replace(",", "").replace("₹", "").strip())
    except ValueError:
        raise ValueError(f"'{v}' is not a number")


def _date(v) -> Optional[str]:
    if v in (None, ""):
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"date '{v}' — use YYYY-MM-DD")


def _pick(value: str, options: list[str], label: str) -> str:
    hit = next((o for o in options if o.lower() == value.strip().lower()), None)
    if not hit:
        raise ValueError(f"unknown {label} '{value}'")
    return hit


async def _lookups() -> dict:
    m = await commission_masters(None)  # type: ignore[arg-type]  # dependency only checks login
    brokers = await db.brokers.find({}, {"_id": 0, "id": 1, "name": 1, "code": 1}).to_list(1000)
    agents = await db.users.find({"role": "agent"}, {"_id": 0, "id": 1, "username": 1, "full_name": 1}).sort("username", 1).to_list(5000)
    groups = [g["name"] for g in await db.agent_groups.find({}, {"_id": 0, "name": 1}).to_list(200)]
    used = [n for n in await db.commission_rules.distinct("insurer") if isinstance(n, str) and n not in m.insurers]
    zones = await db.company_zones.find({}, {"_id": 0, "name": 1, "insurer": 1, "rtos": 1}).sort([("insurer", 1), ("name", 1)]).to_list(2000)
    return {"insurers": [*m.insurers, *used], "brokers": brokers, "agents": agents, "groups": groups, "zones": zones}


def _type_of(products: list[str]) -> str:
    types = {t for t, items in PRODUCT_TYPES.items() for p in products if p in items}
    return types.pop() if len(types) == 1 else ""


def agent_label(a: dict) -> str:
    """'AGT1022 · Sambu' — what the admin sees / picks in the Agent column."""
    return f"{a['username'].upper()} · {a['full_name']}" if a.get("full_name") else a["username"].upper()


def _find_agent(value: str, agents: list[dict]) -> Optional[dict]:
    """Accepts 'AGT1022 · Sambu', 'AGT1022 - Sambu', 'AGT1022' or the agent's full name."""
    v = value.strip().lower()
    code = v.replace(" - ", " · ").split("·")[0].strip()
    return next((a for a in agents if a["username"].lower() == code), None) or \
        next((a for a in agents if (a.get("full_name") or "").strip().lower() == v), None)


def _rule_row(r: dict, bname: dict, acode: dict, cols: list[tuple[str, str]] = COLS) -> list:
    ptype = _type_of(r.get("products") or [])
    products = r.get("products") or []
    out = {
        "id": r["id"], "kind": r["kind"].title(), "name": r["name"], "insurer": r.get("insurer"), "broker": ", ".join(bname.get(b, "") for b in r.get("broker_ids") or []),
        "agent_group": r.get("agent_group"), "agent_code": acode.get(r.get("agent_id") or ""), "ptype": ptype,
        "products": "" if ptype and set(products) == set(PRODUCT_TYPES[ptype]) else ", ".join(products),
        "makes": ", ".join(r.get("makes") or []), "model": r.get("model"), "rtos": ", ".join(r.get("rtos") or []),
        "vehicle_ages": ", ".join(r.get("vehicle_ages") or []), "fuel_types": ", ".join(r.get("fuel_types") or []), "business_types": ", ".join(r.get("business_types") or []),
        "ncb": NCB_OUT.get(r.get("ncb") or "any"), "cover": COVER_OUT.get(r.get("cover") or "any"),
        "valid_from": r.get("valid_from"), "valid_to": r.get("valid_to"), "active": "Yes" if r.get("active", True) else "No",
        **{k: r.get(k) for k in NUMS},
    }
    return [out.get(k) for _, k in cols]


@router.get("/commission/grid/template")
async def grid_template(prefill: bool = False, kind: str = "", _: User = Depends(admin_only)):
    lk = await _lookups()
    wb = Workbook()
    wb.remove(wb.active)
    rules_by_sheet: dict[str, list[dict]] = {}
    if prefill:
        q: dict = {"superseded_by": None}
        if kind in ("receivable", "payable"):
            q["kind"] = kind
        for r in await db.commission_rules.find(q, {"_id": 0}).sort([("kind", -1), ("insurer", 1), ("name", 1)]).to_list(5000):
            r = CommissionRule(**r).model_dump()
            rules_by_sheet.setdefault(_type_of(r.get("products") or []) or ALL_SHEET, []).append(r)

    # Lists sheet: every dropdown's values (per product where they differ)
    lists = wb.create_sheet("Lists")
    columns: dict[str, list] = {
        "Type": ["Receivable", "Payable"], "Insurance Company": lk["insurers"], "Channel Partner": [b["name"] for b in lk["brokers"]],
        "Agent Group": lk["groups"], "Agent Code": [agent_label(a) for a in lk["agents"]], "Product Type": list(PRODUCT_TYPES),
        "Products": ALL_PRODUCTS, "Business Types": BUSINESS_TYPES, "Vehicle Age Bands": AGE_BANDS, "RTOs": rto_options(lk["zones"]),
        "Fuel Types": sorted({f for v in FUEL_BY_TYPE.values() for f in v}), "Makes": sorted({m for v in MAKES.values() for m in v}),
        "NCB": list(NCB_OUT.values()), "Package / SATP": list(COVER_OUT.values()), "Active": ["Yes", "No"],
        **{f"Products · {t}": v for t, v in PRODUCT_TYPES.items()}, **{f"Makes · {t}": v for t, v in MAKES.items()},
        **{f"Fuel · {t}": v for t, v in FUEL_BY_TYPE.items()},
    }
    ranges = {}
    for ci, (title, values) in enumerate(columns.items(), start=1):
        lists.cell(1, ci, title).font = Font(bold=True)
        for ri, v in enumerate(values, start=2):
            lists.cell(ri, ci, v)
        col = lists.cell(1, ci).column_letter
        lists.column_dimensions[col].width = 26
        ranges[title] = f"Lists!${col}$2:${col}${max(2, len(values) + 1)}"

    bname = {b["id"]: b["name"] for b in lk["brokers"]}
    acode = {a["id"]: agent_label(a) for a in lk["agents"]}
    head_fill = PatternFill("solid", fgColor="102A43")
    for idx, sheet in enumerate([*PRODUCT_TYPES, ALL_SHEET]):
        cols = sheet_cols(sheet)
        ws = wb.create_sheet(sheet, idx)
        ws.append([h for h, _ in cols])
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head_fill, Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 42
        ws.freeze_panes = "D2"
        for i, (h, _) in enumerate(cols, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(34, len(h) + 2))
        for r in rules_by_sheet.get(sheet, []):
            ws.append(_rule_row(r, bname, acode, cols))
        typed = sheet in PRODUCT_TYPES
        dropdowns = {
            "kind": "Type", "insurer": "Insurance Company", "broker": "Channel Partner", "agent_group": "Agent Group", "agent_code": "Agent Code",
            "ptype": "Product Type", "products": f"Products · {sheet}" if typed else "Products",
            "makes": f"Makes · {sheet}" if f"Makes · {sheet}" in columns else "Makes",
            "fuel_types": f"Fuel · {sheet}" if f"Fuel · {sheet}" in columns else "Fuel Types",
            "rtos": "RTOs", "vehicle_ages": "Vehicle Age Bands", "business_types": "Business Types",
            "ncb": "NCB", "cover": "Package / SATP", "active": "Active",
        }
        for i, (_, key) in enumerate(cols, start=1):
            if key not in dropdowns:
                continue
            col = ws.cell(1, i).column_letter
            strict = key in ("kind", "ptype", "ncb", "cover", "active")
            dv = DataValidation(type="list", formula1=ranges[dropdowns[key]], allow_blank=True, showErrorMessage=strict)
            if key in MULTI:
                dv.promptTitle, dv.prompt, dv.showInputMessage = "Pick or type", "Pick one, or type several separated by commas", True
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}2000")

    info = wb.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 140

    buf = io.BytesIO()
    neutralise_formulas(wb)
    wb.save(buf)
    buf.seek(0)
    name = f"commission_grid_{kind or 'all'}_{_today()}.xlsx" if prefill else "commission_grid_template.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


def _auto_name(d: dict, ptype: str) -> str:
    parts = [(d.get("insurer") or "All companies").split(" /")[0], ptype or "All products"]
    if d["products"] and ptype and set(d["products"]) != set(PRODUCT_TYPES.get(ptype, [])):
        parts.append(", ".join(d["products"]))
    if d.get("rtos"):
        parts.append(", ".join(d["rtos"][:4]) + ("…" if len(d["rtos"]) > 4 else ""))
    if d.get("makes"):
        parts.append(", ".join(d["makes"]))
    return " · ".join(parts)


def _parse_row(raw: dict, lk: dict, sheet_type: str = "") -> tuple[Optional[str], dict]:
    """Row cells → (rule id or None, CommissionRuleIn kwargs). Raises ValueError with a readable message."""
    s = lambda k: str(raw.get(k) if raw.get(k) is not None else "").strip()  # noqa: E731
    kind = s("kind").lower()
    if kind.startswith("rec"):
        kind = "receivable"
    elif kind.startswith("pay"):
        kind = "payable"
    else:
        raise ValueError("Type must be Receivable or Payable")
    d: dict = {"kind": kind}
    d["insurer"] = _pick(s("insurer"), lk["insurers"], "insurance company") if s("insurer") else None
    d["broker_ids"] = []
    for name in _split(raw.get("broker")):
        b = next((b for b in lk["brokers"] if name.lower() in (b["name"].lower(), (b.get("code") or "").lower())), None)
        if not b:
            raise ValueError(f"unknown channel partner '{name}' — pick from the dropdown")
        d["broker_ids"].append(b["id"])
    if kind == "payable":
        d["agent_group"] = _pick(s("agent_group"), lk["groups"], "agent group") if s("agent_group") else None
        if s("agent_code"):
            a = _find_agent(s("agent_code"), lk["agents"])
            if not a:
                raise ValueError(f"unknown agent '{s('agent_code')}' — pick Code · Name from the dropdown")
            d["agent_id"] = a["id"]
    ptype = _pick(s("ptype"), list(PRODUCT_TYPES), "product type") if s("ptype") else sheet_type
    products = [_pick(p, ALL_PRODUCTS, "product") for p in _split(raw.get("products"))]
    if ptype and products and any(p not in PRODUCT_TYPES[ptype] for p in products):
        raise ValueError(f"products must belong to {ptype}")
    if ptype == "PCV" and not products:
        raise ValueError("PCV rules need Products (e.g. PCV - Taxi)")
    d["products"] = products or (list(PRODUCT_TYPES[ptype]) if ptype else [])
    ptype = ptype or _type_of(products)
    known_makes = MAKES.get(ptype) or sorted({m for v in MAKES.values() for m in v})
    d["makes"] = [next((k for k in known_makes if k.lower() == m.lower()), m) for m in _split(raw.get("makes"))]
    d["model"] = s("model") or None
    d["rtos"] = _rtos(raw.get("rtos"), lk.get("zones", []))
    d["vehicle_ages"] = [_pick(a, AGE_BANDS, "vehicle age band") for a in _split(raw.get("vehicle_ages"))]
    fuels = FUEL_BY_TYPE.get(ptype) or sorted({*FUEL_TYPES, *[f for v in FUEL_BY_TYPE.values() for f in v]})
    d["business_types"] = [_pick(b, BUSINESS_TYPES, "business type") for b in _split(raw.get("business_types"))]
    d["fuel_types"] = [_pick(f, fuels, f"fuel type for {ptype or 'this product'}") for f in _split(raw.get("fuel_types"))]
    ncb = s("ncb").lower()
    d["ncb"] = "with" if ncb.startswith("with ") or ncb == "with" else "without" if ncb.startswith("without") else "any"
    cover = s("cover").lower()
    d["cover"] = "package" if cover.startswith("pack") else "satp" if cover.startswith("satp") or cover.startswith("liab") else "any"
    for k in NUMS:
        d[k] = _num(raw.get(k))
    d["valid_from"], d["valid_to"] = _date(raw.get("valid_from")), _date(raw.get("valid_to"))
    d["active"] = s("active").lower() not in ("no", "n", "false", "0", "inactive")
    d["name"] = s("name") or _auto_name(d, ptype)
    return (s("id") or None), d


def _same_rate(old: dict, d: dict) -> bool:
    return all(old.get(k) == d.get(k) for k in (*COMPONENTS, "valid_to", "active", "name"))


def _fmt(v) -> str:
    return "—" if v is None else f"{v:g}"


def _rates_text(d: dict) -> list[str]:
    return [f"{LABEL[c]} {_fmt(d[c])}" for c in COMPONENTS if d.get(c) is not None]


COND_LABELS = {"insurer": "Company", "broker_ids": "Channel partners", "agent_group": "Agent group", "agent_id": "Agent", "products": "Products", "makes": "Makes",
               "model": "Model", "rtos": "RTOs", "vehicle_ages": "Vehicle age", "fuel_types": "Fuel", "business_types": "Business type", "ncb": "NCB", "cover": "Package/SATP",
               "gvw_min": "GVW from", "gvw_max": "GVW to", "cc_min": "CC from", "cc_max": "CC to", "seating_min": "Seating from", "seating_max": "Seating to"}


def _cond_diff(old: dict, d: dict, names: dict) -> list[str]:
    """Each changed matching condition as 'RTOs TN19 → TN19, TN20'."""
    def show(k, v):
        if k == "broker_ids":
            return ", ".join(sorted(names.get(b, b) for b in v or [])) or "all"
        if isinstance(v, list):
            return ", ".join(sorted(v)) or "all"
        if k == "agent_id":
            return names.get(v or "", v) or "all"
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return f"{v:g}"
        return str(v) if v not in (None, "") else ("any" if k in ("ncb", "cover") else "all")
    out = []
    for k, label in COND_LABELS.items():
        a, b = old.get(k), d.get(k)
        if isinstance(a, list) or isinstance(b, list):
            if sorted(a or []) == sorted(b or []):
                continue
        elif show(k, a).lower() == show(k, b).lower():
            continue
        out.append(f"{label}: {show(k, a)} → {show(k, b)}")
    return out


def _diff(old: dict, d: dict, same_conditions: bool, names: dict | None = None) -> list[str]:
    out = [f"{LABEL[c]} {_fmt(old.get(c))} → {_fmt(d.get(c))}" for c in COMPONENTS if old.get(c) != d.get(c)]
    if old.get("name") != d.get("name"):
        out.append(f"Name '{old.get('name')}' → '{d.get('name')}'")
    if old.get("active", True) != d.get("active", True):
        out.append("Activated" if d.get("active", True) else "Deactivated")
    if old.get("valid_to") != d.get("valid_to"):
        out.append(f"Applicable to {old.get('valid_to') or 'open'} → {d.get('valid_to') or 'open'}")
    if not same_conditions:
        out += _cond_diff(old, d, names or {}) or ["Conditions changed"]
    return out


def _apply_dates(d: dict, valid_from: Optional[str], valid_to: Optional[str], apply_to_past: bool) -> None:
    """The check screen's From / To dates override the row's own dates — only for rows that change."""
    if valid_from and not apply_to_past:
        d["valid_from"] = valid_from
    if valid_to:
        d["valid_to"] = valid_to
    if d.get("valid_from") and d.get("valid_to") and d["valid_to"] < d["valid_from"]:
        raise ValueError(f"Applicable to {d['valid_to']} is before applicable from {d['valid_from']} — set a To date on the check screen")


def _sim_validate(rule: dict, others: list[dict]) -> None:
    """Same check as saving (payable ≤ overlapping receivable), against the simulated grid."""
    if not rule["active"]:
        return
    others = [o for o in others if o.get("active", True)]
    if rule["kind"] == "payable":
        err = check_against(rule, [r for r in others if r["kind"] == "receivable"])
    else:
        err = next((e for p in others if p["kind"] == "payable" and (e := check_against(p, [rule]))), None)
        err = err and f"{err} — lower the payable rule first"
    if err:
        raise HTTPException(status_code=400, detail=err)


@router.post("/commission/grid/upload", response_model=GridUploadResult)
async def grid_upload(file: UploadFile = File(...), preview: bool = False, valid_from: Optional[str] = None, valid_to: Optional[str] = None,
                      apply_to_past: bool = False, _: User = Depends(admin_only)):
    """preview=true → check every row and report what would change, saving nothing.
    valid_from / valid_to (optional) override every row's Applicable From / To. apply_to_past → changed rules are corrected for every
    policy (older versions replaced) and new rules apply to all policies."""
    try:
        valid_from, valid_to = _date(valid_from), _date(valid_to)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    if valid_from and valid_to and valid_to < valid_from:
        raise HTTPException(status_code=400, detail="'To' date must be on or after the 'From' date")
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload the .xlsx grid template")
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file")
    by_header = {h.lower(): k for h, k in COLS}
    by_header["agent code (payable)"] = "agent_code"  # older downloaded files
    by_header["channel partner"] = "broker"
    parsed: list[tuple[int, dict, str]] = []  # (row number, cells, sheet)
    found = False
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        rows = list(ws.iter_rows(values_only=True))
        keys = [by_header.get(str(h or "").strip().lower()) for h in rows[0]] if rows else []
        if "kind" not in keys or not any(k in keys for k in COMPONENTS):
            continue
        found = True
        for n, row in enumerate(rows[1:], start=2):
            raw = {k: v for k, v in zip(keys, row) if k}
            if not all(v in (None, "") for v in raw.values()):
                parsed.append((n, raw, ws.title))
    if not found:
        raise HTTPException(status_code=400, detail="Header row not recognised — download the template and keep its first row")
    return await process_rows(parsed, preview, valid_from, valid_to, apply_to_past)


async def process_rows(parsed: list[tuple[int, dict, str]], preview: bool, valid_from: Optional[str] = None, valid_to: Optional[str] = None,
                       apply_to_past: bool = False) -> GridUploadResult:
    """Check (preview) or save grid rows given as (row number, {column key: value}, sheet) — shared by Excel and AI uploads."""
    lk = await _lookups()
    names = {b["id"]: b["name"] for b in lk["brokers"]} | {a["id"]: agent_label(a) for a in lk["agents"]}
    res = GridUploadResult(preview=preview)
    parsed.sort(key=lambda nr: 0 if str(nr[1].get("kind") or "").lower().startswith("rec") else 1)  # receivables first

    # normalised through the model so older rules (missing newer fields) compare equal to a re-uploaded row
    current = [CommissionRule(**r).model_dump(exclude={"apply_to_past"}) for r in await db.commission_rules.find({"superseded_by": None}, {"_id": 0}).to_list(5000)]
    newer = {r["id"]: r["superseded_by"] for r in await db.commission_rules.find({"superseded_by": {"$nin": [None, "ended"]}}, {"_id": 0, "id": 1, "superseded_by": 1}).to_list(10000)}
    sim_all = await db.commission_rules.find({"active": True}, {"_id": 0}).to_list(10000) if preview else []
    ident = lambda r: (r["kind"], r.get("agent_group"), r.get("agent_id"), _signature(r))  # noqa: E731
    for n, raw, sheet in parsed:
        kind_hint = str(raw.get("kind") or "").strip().title() or None
        msg = "Row could not be read"
        try:
            rid, kw = _parse_row(raw, lk, sheet if sheet in PRODUCT_TYPES else "")
            d = _clean(CommissionRuleIn(**kw))  # the screen's To date is applied below, only to rows that change
            for _ in range(50):  # a Rule ID from an older download → follow to the current version
                if rid not in newer:
                    break
                rid = newer[rid]
            old = next((r for r in current if r["id"] == rid), None) if rid else next((r for r in current if ident(r) == ident(d)), None)
            if rid and not old:
                raise ValueError(f"Rule ID {rid} not found (or it is an old version) — clear the Rule ID to add it as a new rule")
            info = {"row": n, "sheet": sheet, "kind": d["kind"], "name": d["name"], "insurer": d.get("insurer")}
            if old:
                if old["kind"] != d["kind"]:
                    raise ValueError("Type cannot change for an existing Rule ID")
                same_conditions = ident(old) == ident(d)
                if same_conditions and _same_rate(old, d) and (apply_to_past or valid_from or not d["valid_from"] or d["valid_from"] == old.get("valid_from")):
                    res.unchanged += 1
                    res.rows.append(GridRowPreview(**info, action="unchanged"))
                    continue
                _apply_dates(d, valid_from, valid_to, apply_to_past)
                if apply_to_past:
                    d["valid_from"] = old.get("valid_from")  # keep the rule's own start; the change covers all its policies
                    eff = "All policies"
                else:
                    if d["valid_from"] == old.get("valid_from") and not valid_from and not _same_rate(old, d):
                        d["valid_from"] = None  # only the rate changed → new rate from today, past policies keep the old one
                    eff = d["valid_from"] or _today()
                    if old.get("valid_from") and eff <= old["valid_from"]:
                        eff = old["valid_from"]
                changes = _diff(old, d, same_conditions, names)
                if preview:
                    _sim_validate({**d, "id": old["id"]}, [r for r in sim_all if r["id"] != old["id"]])
                    sim_all = [r for r in sim_all if r["id"] != old["id"]] + [{**d, "id": old["id"]}]
                    current = [r for r in current if r["id"] != old["id"]] + [{**d, "id": old["id"]}]
                else:
                    saved = await apply_rule_update(old, d, apply_to_past)
                    current = [r for r in current if r["id"] != old["id"]] + [saved.model_dump()]
                    if saved.id != old["id"]:
                        newer[old["id"]] = saved.id
                res.updated += 1
                res.rows.append(GridRowPreview(**info, action="update", changes=changes, effective_from=eff))
            else:
                _apply_dates(d, valid_from, valid_to, apply_to_past)
                if apply_to_past:
                    d["valid_from"] = None
                if preview:
                    _sim_validate(d, sim_all)
                    sim_all.append({**d, "id": f"row-{sheet}-{n}"})
                    current.append({**d, "id": f"row-{sheet}-{n}"})
                else:
                    current.append((await insert_rule(d)).model_dump())
                res.created += 1
                res.rows.append(GridRowPreview(**info, action="create", changes=_rates_text(d), effective_from="All policies" if apply_to_past else d["valid_from"]))
            continue
        except HTTPException as e:
            msg = str(e.detail)
        except ValidationError as e:
            err = e.errors()[0]
            msg = f"{'.'.join(str(x) for x in err['loc'])}: {err['msg']}"
        except ValueError as e:
            msg = str(e)
        res.errors.append(GridRowError(row=n, sheet=sheet, error=msg))
        res.rows.append(GridRowPreview(row=n, sheet=sheet, action="error", kind=kind_hint, name=str(raw.get("name") or "") or None, insurer=str(raw.get("insurer") or "") or None, error=msg))
    order = {t: i for i, t in enumerate([*PRODUCT_TYPES, ALL_SHEET])}
    res.rows.sort(key=lambda r: (order.get(r.sheet or "", 99), r.sheet or "", r.row))
    return res


# ---------- receivable vs payable comparison ----------
def _rates(d: dict) -> CompareRates:
    return CompareRates(**{c: d.get(c) for c in COMPONENTS})


def _margin(rec: dict, pay: dict) -> CompareRates:
    return CompareRates(**{c: round((rec.get(c) or 0) - pay[c], 4) for c in COMPONENTS if pay.get(c) is not None})


def _conditions(r: dict) -> list[str]:
    out = []
    if r.get("rtos"):
        out.append("RTO " + ", ".join(r["rtos"][:6]) + ("…" if len(r["rtos"]) > 6 else ""))
    if r.get("makes"):
        out.append(", ".join(r["makes"]) + (f" {r['model']}" if r.get("model") else ""))
    if r.get("vehicle_ages"):
        out.append(", ".join(r["vehicle_ages"]))
    if r.get("fuel_types"):
        out.append("/".join(r["fuel_types"]))
    if r.get("business_types"):
        out.append("/".join(r["business_types"]))
    if (r.get("cover") or "any") != "any":
        out.append(r["cover"].upper())
    if (r.get("ncb") or "any") != "any":
        out.append(f"{r['ncb']} NCB")
    for dim, unit in (("gvw", "kg"), ("cc", "CC"), ("seating", "seats")):
        if r.get(f"{dim}_min") is not None or r.get(f"{dim}_max") is not None:
            out.append(f"{_fmt(r.get(f'{dim}_min') or 0)}–{_fmt(r.get(f'{dim}_max')) if r.get(f'{dim}_max') is not None else '∞'} {unit}")
    return out


@router.get("/commission/compare", response_model=List[CompareRow])
async def grid_compare(_: User = Depends(admin_only)):
    return await compare_rows()


def _rates_cell(r: CompareRates, margin: bool = False) -> str:
    parts = []
    for k, label in (("od_percent", "OD"), ("tp_percent", "TP"), ("net_percent", "Net"), ("odtp_percent", "OD+TP"), ("flat_amount", "Flat ₹")):
        v = getattr(r, k)
        if v is not None:
            parts.append(f"{label} {'+' if margin and v > 0 else ''}{v:g}{'' if k == 'flat_amount' else '%'}")
    return ", ".join(parts) or "—"


@router.get("/commission/compare/export")
async def grid_compare_export(insurer: str = "", _: User = Depends(admin_only)):
    """Compare & margin view as Excel — one sheet per insurance company (or just the one asked for)."""
    rows = [r for r in await compare_rows() if not insurer or (r.insurer or "All companies") == insurer]
    if not rows:
        raise HTTPException(status_code=404, detail="No receivable rules for this company")
    wb = Workbook()
    wb.remove(wb.active)
    head = ["Product", "Channel Partner", "Receivable Rule", "Conditions", "Applicable From", "Receivable", "Paid To", "Payable Source", "Payable", "Your Margin",
            *[f"{label} Margin" for label in ("OD %", "TP %", "Net %", "OD+TP %", "Flat ₹")]]
    fills = {"pos": PatternFill("solid", fgColor="D1FAE5"), "zero": PatternFill("solid", fgColor="FEF3C7"), "neg": PatternFill("solid", fgColor="FEE2E2")}
    used: set[str] = set()
    for company in dict.fromkeys(r.insurer or "All companies" for r in rows):
        title = "".join(ch for ch in company if ch not in '[]:*?/\\')[:31] or "Sheet"
        while title in used:
            title = title[:28] + f"_{len(used)}"
        used.add(title)
        ws = wb.create_sheet(title)
        ws.append([f"{company} — receivable vs payable (margin = receivable − payable)"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(head)
        for c in ws[2]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0F766E"), Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "B3"
        for r in (x for x in rows if (x.insurer or "All companies") == company):
            base = [r.product, r.broker or "", r.rule_name, " · ".join(r.conditions), r.valid_from or "Always", _rates_cell(r.rates)]
            if not r.payables:
                ws.append([*base, "NO PAYABLE SET", "", "", ""])
                ws.cell(ws.max_row, 7).fill = fills["zero"]
                continue
            for p in r.payables:
                ws.append([*base, p.label, "Rule: " + (p.rule_name or "") if p.source == "rule" else "Group margin", _rates_cell(p.rates), _rates_cell(p.margin, True),
                           *[getattr(p.margin, k) for k in COMPONENTS]])
                vals = [v for v in (getattr(p.margin, k) for k in COMPONENTS) if v is not None]
                if vals:
                    ws.cell(ws.max_row, 10).fill = fills["neg" if min(vals) < 0 else "zero" if p.low_margin or min(vals) == 0 else "pos"]
                if p.low_margin:
                    ws.cell(ws.max_row, 10).value = f"LOW · {ws.cell(ws.max_row, 10).value}"
        for col, w in zip("ABCDEFGHIJKLMNO", (26, 22, 30, 36, 14, 20, 26, 26, 20, 22, 10, 10, 10, 10, 10)):
            ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    neutralise_formulas(wb)
    wb.save(buf)
    buf.seek(0)
    tag = "".join(ch if ch.isalnum() else "_" for ch in insurer)[:30] if insurer else "all_companies"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="commission_margin_{tag}_{_today()}.xlsx"'})


def _low(m: CompareRates, min_margin: float) -> bool:
    return any(getattr(m, k) is not None and getattr(m, k) < min_margin for k in PCT_KEYS)


async def get_settings() -> CommissionSettings:
    d = await db.app_settings.find_one({"key": "commission"}, {"_id": 0})
    return CommissionSettings(**(d or {}).get("value", {}))


@router.get("/commission/settings", response_model=CommissionSettings)
async def read_settings(_: User = Depends(admin_only)):
    return await get_settings()


@router.put("/commission/settings", response_model=CommissionSettings)
async def save_settings(body: CommissionSettings, _: User = Depends(admin_only)):
    await db.app_settings.update_one({"key": "commission"}, {"$set": {"value": body.model_dump()}}, upsert=True)
    return body


@router.post("/commission/compare/bulk-payable", response_model=BulkPayableResult)
async def bulk_payable(body: BulkPayableIn, _: User = Depends(admin_only)):
    """One agent group's payable on every receivable with no payable yet (payable % = receivable % − margin points)."""
    group = await db.agent_groups.find_one({"name": body.agent_group}, {"_id": 0, "name": 1})
    if not group:
        raise HTTPException(status_code=404, detail="Agent group not found")
    targets = [r for r in await compare_rows() if not r.payables and (not body.insurer or (r.insurer or "All companies") == body.insurer)
               and (body.rule_ids is None or r.rule_id in body.rule_ids)]
    recs = {r["id"]: r for r in await db.commission_rules.find({"id": {"$in": [t.rule_id for t in targets]}}, {"_id": 0}).to_list(5000)}
    res = BulkPayableResult(preview=body.preview)
    for t in targets:
        rec = recs.get(t.rule_id)
        pay = derive_payable(rec, body.margin) if rec else {}
        rates = CompareRates(**{k: pay.get(k) for k in PCT_KEYS})
        row = BulkPayableRow(rule_id=t.rule_id, insurer=t.insurer, product=t.product, receivable_name=t.rule_name, receivable=t.rates, payable=rates, status="ready")
        if not rec or all(pay.get(k) is None for k in PCT_KEYS):
            row.status, row.note = "skipped", "Receivable has only a flat amount — set this payable by hand"
        elif not body.preview:
            keep = ("insurer", "products", "rtos", "vehicle_ages", "fuel_types", "business_types", "ncb", "cover", "broker_ids", "gvw_min", "gvw_max", "cc_min", "cc_max",
                    "seating_min", "seating_max", "makes", "model", "valid_from", "valid_to")
            try:
                d = _clean(CommissionRuleIn(kind="payable", name=f"{rec['name']} · Group {group['name']}", agent_group=group["name"],
                                            **{k: rec.get(k) for k in keep if rec.get(k) is not None}, **{k: pay.get(k) for k in PCT_KEYS}))
                await insert_rule(d)
                row.status = "created"
                res.created += 1
            except HTTPException as e:
                row.status, row.note = "error", str(e.detail)
            except ValidationError as e:
                row.status, row.note = "error", e.errors()[0]["msg"]
        res.rows.append(row)
    return res


async def compare_rows() -> List[CompareRow]:
    """Every current receivable rule beside the payables that overlap it (explicit rules + group margins), with your margin."""
    today = _today()
    min_margin = (await get_settings()).min_margin
    live = lambda r: r.get("active", True) and not r.get("superseded_by") and (not r.get("valid_to") or r["valid_to"] >= today)  # noqa: E731
    rules = [r for r in await db.commission_rules.find({}, {"_id": 0}).to_list(10000) if live(r)]
    brokers = {b["id"]: b["name"] for b in await db.brokers.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(1000)}
    agents = {a["id"]: f"{a.get('full_name') or a['username']} ({a['username'].upper()})" for a in await db.users.find({"role": "agent"}, {"_id": 0, "id": 1, "username": 1, "full_name": 1}).to_list(5000)}
    groups = await db.agent_groups.find({}, {"_id": 0}).sort("name", 1).to_list(200)
    recs = sorted([r for r in rules if r["kind"] == "receivable"], key=lambda r: ((r.get("insurer") or "~").lower(), product_type_of(r.get("products") or []) or "~", r["name"]))
    pays = [r for r in rules if r["kind"] == "payable"]
    out = []
    for rec in recs:
        ptype = product_type_of(rec.get("products") or [])
        prods = rec.get("products") or []
        product = "All products" if not prods else (f"{ptype} (all)" if ptype and set(prods) == set(PRODUCT_TYPES.get(ptype, [])) else ", ".join(prods))
        row = CompareRow(rule_id=rec["id"], insurer=rec.get("insurer"), product=product, broker=", ".join(brokers.get(b, "") for b in rec.get("broker_ids") or []) or None, rule_name=rec["name"],
                         conditions=_conditions(rec), valid_from=rec.get("valid_from"), rates=_rates(rec))
        for p in pays:
            if overlaps(p, rec):
                who = agents.get(p.get("agent_id") or "") or (f"Group {p['agent_group']}" if p.get("agent_group") else "All agents")
                mg = _margin(rec, p)
                row.payables.append(ComparePayable(label=who, source="rule", rule_name=p["name"], rates=_rates(p), margin=mg, low_margin=_low(mg, min_margin)))
        for g in groups:
            line = group_margin_line(g.get("margins") or [], rec.get("insurer"), None, ptype, today, single_product(prods))
            if line is not None:
                derived = derive_payable(rec, line)
                mg = _margin(rec, derived)
                row.payables.append(ComparePayable(label=f"Group {g['name']} ({margin_text(line)})", source="group_margin", rates=_rates(derived), margin=mg, low_margin=_low(mg, min_margin)))
        out.append(row)
    return out


# ---------- AI grid reading: Excel / PDF / image → rules to approve ----------
AI_TYPES = {".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12"}
AI_MAX_BYTES = 15 * 1024 * 1024


@router.post("/commission/grid/ai-read", response_model=AiGridRead)
async def grid_ai_read(file: UploadFile = File(...), kind: str = "receivable", broker_id: str = "", insurer: str = "", _: User = Depends(admin_only)):
    """AI reads any commission grid file; nothing is saved — the rows come back for review."""
    import os
    import tempfile
    import uuid
    from lib.grid_ai import read_grid
    name = file.filename or "grid"
    ext = os.path.splitext(name.lower())[1]
    if ext not in AI_TYPES:
        raise HTTPException(status_code=400, detail="Upload an Excel (.xlsx), PDF or image (PNG / JPG / WEBP) grid")
    data = await file.read()
    if len(data) > AI_MAX_BYTES:
        raise HTTPException(status_code=400, detail="File is larger than 15 MB")
    lk = await _lookups()
    hints = {"insurers": lk["insurers"], "brokers": [b["name"] for b in lk["brokers"]], "ptypes": list(PRODUCT_TYPES),
             "products": ALL_PRODUCTS, "ages": AGE_BANDS}
    partner = next((b["name"] for b in lk["brokers"] if b["id"] == broker_id), None) if broker_id else None
    if broker_id and not partner:
        raise HTTPException(status_code=400, detail="Unknown channel partner")
    hints["broker"] = partner
    kind = kind if kind in ("receivable", "payable") else "receivable"
    template = await find_template(kind, insurer or None, name)
    if template:
        hints["template"] = prompt_block(template)
    if insurer:
        hints["insurers"] = [insurer]
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(data)
        path = tmp.name
    try:
        rows, note = await read_grid(path, AI_TYPES[ext], name, kind, hints, f"grid-{uuid.uuid4()}")
    except (RuntimeError, ValueError) as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception:
        raise HTTPException(status_code=502, detail="The AI could not read this file — try a clearer scan or the Excel template")
    finally:
        os.unlink(path)
    if not rows:
        raise HTTPException(status_code=422, detail="No commission rates found in this file")
    for r in rows:
        if partner:  # each channel partner sends its own grid — the picked partner applies to every row
            r.broker = partner
        if insurer:  # picked company applies to every row
            r.insurer = insurer
    return AiGridRead(file_name=name, rows=rows, note=note or None,
                      template_insurer=template["insurer"] if template else None, template_uses=template.get("uses", 0) if template else 0, today=_today())


@router.post("/commission/grid/ai-apply", response_model=GridUploadResult)
async def grid_ai_apply(body: AiGridApplyIn, user: User = Depends(admin_only)):
    """The rows the admin approved (and edited) → same checks / save as the Excel upload."""
    try:
        valid_from, valid_to = _date(body.valid_from), _date(body.valid_to)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    parsed = [(i, {k: v for k, v in r.model_dump(exclude={"needs_review", "review_note"}).items() if v not in (None, "")}, "AI")
              for i, r in enumerate(body.rows, start=1)]
    res = await process_rows(parsed, body.preview, valid_from, valid_to)
    if not body.preview and not res.errors:  # remember this company's layout + the admin's corrections for next time
        kind = (body.rows[0].kind or "Receivable").lower()
        res.learned_insurer, res.learned_accuracy = await learn("payable" if kind.startswith("pay") else "receivable", body.template_insurer, body.rows,
                                                                body.original_rows, body.file_name, user, body.used_layout)
    return res
