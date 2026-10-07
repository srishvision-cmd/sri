import asyncio
import csv
import os
import io
import logging
import re
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import List, Literal, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response, StreamingResponse
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from lib.auth import current_user, require_roles
from lib.dates import today_iso
from lib.db import db
from lib.compress import TARGET_BYTES, compress_document
from lib.commission import commissions_for
from lib.ocr import MODEL_NAME, extract_policy
from lib.storage import get_object, object_path, put_object
from models.common import now_utc
from models.policy import (
    AGE_BUCKETS,
    BRANCHES,
    OUTSTANDING_MODES,
    PAYMENT_MODES,
    PRODUCT_CATEGORIES,
    AgentOutstanding,
    PaymentEntry,
    PaymentIn,
    PaymentModeIn,
    BrokerSetIn,
    BulkBrokerIn,
    ReassignAgentIn,
    ReassignResult,
    payment_totals,
    COVER_TYPES,
    GVW_NA_CLASSES,
    VEHICLE_CLASSES,
    DERIVED_FIELDS,
    derive_fields,
    LOBS,
    FlagIn,
    GroupStat,
    HistoryEntry,
    Meta,
    Policy,
    RENEWAL_GAP_MONTHS,
    PolicyFields,
    PolicyUpdate,
    Stats,
    Usage,
    UsageGroup,
    UsageMonth,
    VerifyIn,
)
from models.user import User
from routers.users import agent_options_for

logger = logging.getLogger(__name__)
router = APIRouter(tags=["policies"])
reviewers = require_roles("admin", "backoffice")

MAX_BYTES = 15 * 1024 * 1024
ALLOWED = {"application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp"}
EXT_TO_MIME = {".pdf": "application/pdf", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}
_tasks: set = set()
# OCR jobs run concurrently on worker threads; cap them so a big batch can't exhaust the pod
OCR_CONCURRENCY = int(os.environ.get("OCR_CONCURRENCY") or 8)
_ocr_slots = asyncio.Semaphore(OCR_CONCURRENCY)


# ---------- scoping ----------
NOT_DELETED = {"is_deleted": {"$ne": True}}


def scope_filter(user: User) -> dict:
    if user.role == "admin":
        return {}
    if user.role == "agent":
        return {"agent_id": user.id}
    ors: list = [{"branch": user.branch}, {"uploaded_by_id": user.id}]
    if user.assigned_agent_ids:
        ors.append({"agent_id": {"$in": user.assigned_agent_ids}})
    return {"$or": ors}


class PolicyFilters:
    def __init__(
        self,
        status: Optional[str] = None,
        lob: Optional[str] = None,
        agent_id: Optional[str] = None,
        branch: Optional[str] = None,
        q: Optional[str] = None,
        date_from: Optional[str] = None,
        date_to: Optional[str] = None,
        age_bucket: Optional[str] = None,
        payment_mode: Optional[str] = None,
    ):
        self.payment_mode = payment_mode
        self.status, self.lob, self.agent_id, self.branch = status, lob, agent_id, branch
        self.age_bucket = age_bucket
        self.q, self.date_from, self.date_to = q, date_from, date_to

    def _equality_clauses(self) -> list:
        pairs = (
            ("status", self.status),
            ("fields.lob", self.lob),
            ("agent_id", self.agent_id),
            ("fields.vehicle_age_bucket", self.age_bucket),
            ("branch", self.branch),
            ("payment_mode", self.payment_mode),
        )
        return [{field: value} for field, value in pairs if value]

    def _search_clause(self) -> Optional[dict]:
        if not self.q:
            return None
        rx = {"$regex": re.escape(self.q.strip()), "$options": "i"}
        keys = ("fields.policy_number", "fields.insured_name", "fields.vehicle_reg_no", "fields.chassis_no", "fields.engine_no", "fields.insurer", "agent_name", "agent_code", "file_name")
        return {"$or": [{k: rx} for k in keys]}

    def _date_clause(self) -> Optional[dict]:
        rng: dict = {}
        try:
            if self.date_from:
                rng["$gte"] = datetime.fromisoformat(self.date_from).replace(tzinfo=timezone.utc)
            if self.date_to:
                rng["$lt"] = datetime.fromisoformat(self.date_to).replace(tzinfo=timezone.utc) + timedelta(days=1)
        except ValueError:
            raise HTTPException(status_code=400, detail="Dates must be YYYY-MM-DD")
        return {"created_at": rng} if rng else None

    def query(self, user: User) -> dict:
        # rejected duplicates never count in MIS, stats or outstanding — shown only when filtered explicitly
        hide_dupes = None if self.status else {"status": {"$ne": "duplicate"}}
        clauses = [NOT_DELETED, hide_dupes, scope_filter(user), *self._equality_clauses(), self._search_clause(), self._date_clause()]
        clauses = [c for c in clauses if c]
        return {"$and": clauses} if clauses else {}


async def load_scoped(policy_id: str, user: User) -> dict:
    query = {"$and": [{"id": policy_id}, NOT_DELETED, scope_filter(user)]}
    doc = await db.policies.find_one(query, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Policy not found")
    return doc


def history(user: User | None, action: str, note: Optional[str] = None) -> dict:
    if user is None:
        return HistoryEntry(by_name="AI OCR Engine", by_role="system", action=action, note=note).model_dump()
    return HistoryEntry(by_name=user.full_name, by_role=user.role, action=action, note=note).model_dump()


# ---------- duplicate check ----------
def _add_months(d: date, months: int) -> date:
    y, m = divmod(d.month - 1 + months, 12)
    y, m = d.year + y, m + 1
    days = [31, 29 if (y % 4 == 0 and (y % 100 != 0 or y % 400 == 0)) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return date(y, m, min(d.day, days))


def _parse_date(v) -> Optional[date]:
    try:
        return date.fromisoformat(str(v)[:10]) if v else None
    except ValueError:
        return None


def _is_renewal(a, b) -> bool:
    """True when the two risk start dates are at least RENEWAL_GAP_MONTHS apart (e.g. 25-10-2025 vs 25-08-2026)."""
    da, db_ = _parse_date(a), _parse_date(b)
    if not da or not db_:
        return False
    early, late = sorted((da, db_))
    return late >= _add_months(early, RENEWAL_GAP_MONTHS)


def _dup_label(d: dict) -> str:
    created = d.get("created_at")
    when = created.strftime("%d-%m-%Y") if isinstance(created, datetime) else ""
    return f"{d.get('agent_name')} ({d.get('agent_code')}) on {when}".strip()


async def find_duplicate(fields: dict, exclude_id: str) -> Optional[tuple[str, str]]:
    """(existing_id, reason) if another live policy has the same policy no, or the same vehicle within the renewal gap."""
    pn = (fields.get("policy_number") or "").strip()
    reg = (fields.get("vehicle_reg_no") or "").strip().upper()
    ors = ([{"fields.policy_number": pn}] if pn else []) + ([{"fields.vehicle_reg_no": reg}] if reg else [])
    if not ors:
        return None
    base = {"id": {"$ne": exclude_id}, "is_deleted": {"$ne": True}, "status": {"$nin": ["duplicate", "processing"]}, "$or": ors}
    proj = {"_id": 0, "id": 1, "fields.policy_number": 1, "fields.vehicle_reg_no": 1, "fields.start_date": 1, "agent_name": 1, "agent_code": 1, "created_at": 1}
    for d in await db.policies.find(base, proj).to_list(50):
        f = d.get("fields") or {}
        if pn and f.get("policy_number") == pn:
            return d["id"], f"Policy no. {pn} already uploaded by {_dup_label(d)}"
    for d in await db.policies.find(base, proj).to_list(50):
        f = d.get("fields") or {}
        if reg and f.get("vehicle_reg_no") == reg and not _is_renewal(f.get("start_date"), fields.get("start_date")):
            return d["id"], (
                f"Vehicle {reg} already has a policy starting {f.get('start_date') or 'unknown'} (uploaded by {_dup_label(d)}); "
                f"a new policy is accepted only if its start date is {RENEWAL_GAP_MONTHS}+ months apart"
            )
    return None


# ---------- brokers ----------
def _bnorm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


async def match_broker(printed: Optional[str]) -> Optional[dict]:
    text = _bnorm(printed)
    if not text:
        return None
    for b in await db.brokers.find({"active": True}, {"_id": 0, "id": 1, "name": 1, "aliases": 1}).to_list(1000):
        if any(n and (n in text or text in n) for n in map(_bnorm, [b["name"], *b.get("aliases", [])])):
            return b
    return None


async def _resolve_broker(broker_id: Optional[str], agent_id: str) -> Optional[dict]:
    """Explicit pick at upload > the agent's default broker > (later) OCR name match."""
    if broker_id:
        b = await db.brokers.find_one({"id": broker_id}, {"_id": 0, "id": 1, "name": 1})
        if not b:
            raise HTTPException(status_code=400, detail="Broker not found")
        return b
    agent = await db.users.find_one({"id": agent_id}, {"_id": 0, "broker_id": 1})
    if agent and agent.get("broker_id"):
        return await db.brokers.find_one({"id": agent["broker_id"]}, {"_id": 0, "id": 1, "name": 1})
    return None


# ---------- OCR ----------
async def _run_ocr(policy_id: str, storage_path: str, mime: str, original: Optional[bytes] = None) -> None:
    async with _ocr_slots:
        await _run_ocr_job(policy_id, storage_path, mime, original)


async def _run_ocr_job(policy_id: str, storage_path: str, mime: str, original: Optional[bytes]) -> None:
    update: dict = {"updated_at": now_utc()}
    try:
        data = original if original is not None else (await get_object(storage_path))[0]
        suffix = ALLOWED.get(mime, ".bin")
        with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:  # OCR SDK reads from a local path
            tmp.write(data)
            tmp.flush()
            fields, confidence = await extract_policy(tmp.name, mime, f"ocr-{policy_id}")
        filled = sum(1 for v in fields.values() if v is not None)
        update.update(status="extracted", fields=fields, confidence=confidence, ocr_error=None, updated_at=now_utc())
        entry = history(None, "OCR extracted", f"{filled} fields captured")
        current = await db.policies.find_one({"id": policy_id}, {"_id": 0, "broker_id": 1})
        if current and not current.get("broker_id"):  # no broker chosen → match the name printed on the policy
            b = await match_broker(fields.get("intermediary_name"))
            if b:
                update.update(broker_id=b["id"], broker_name=b["name"])
        dup = await find_duplicate(fields, policy_id)
        if dup:
            update.update(status="duplicate", duplicate_of_id=dup[0], duplicate_reason=dup[1])
            await db.policies.update_one({"id": policy_id, "status": "processing"}, {"$push": {"history": entry}})
            entry = history(None, "Rejected as duplicate", dup[1])
    except Exception as exc:  # surface any engine failure as a failed record for manual entry
        logger.exception("OCR failed for %s", policy_id)
        update.update(status="failed", ocr_error=str(exc)[:300] or "OCR failed")
        entry = history(None, "OCR failed", update["ocr_error"])
    await db.policies.update_one({"id": policy_id, "status": "processing"}, {"$set": update, "$push": {"history": entry}})


async def requeue_stuck_ocr() -> None:
    """On boot, restart OCR for records left in 'processing' by a previous process (reload/crash)."""
    docs = await db.policies.find(
        {"status": "processing", "is_deleted": {"$ne": True}, "storage_path": {"$exists": True}},
        {"_id": 0, "id": 1, "storage_path": 1, "file_mime": 1},
    ).to_list(500)
    for d in docs:
        _spawn_ocr(d["id"], d["storage_path"], d["file_mime"])
    if docs:
        logger.info("Re-queued OCR for %d policies", len(docs))


def _spawn_ocr(policy_id: str, storage_path: str, mime: str, original: Optional[bytes] = None, original_mime: Optional[str] = None) -> None:
    """OCR reads the uncompressed upload when available (best accuracy); storage keeps the compressed copy."""
    task = asyncio.create_task(_run_ocr(policy_id, storage_path, original_mime or mime, original))
    _tasks.add(task)
    task.add_done_callback(_tasks.discard)


# ---------- routes ----------
@router.get("/meta", response_model=Meta)
async def meta():
    return Meta(branches=BRANCHES, lobs=LOBS, age_buckets=AGE_BUCKETS, cover_types=COVER_TYPES, vehicle_classes=VEHICLE_CLASSES, payment_modes=PAYMENT_MODES, product_categories=PRODUCT_CATEGORIES)


@router.post("/policies/upload", response_model=Policy)
async def upload_policy(
    file: UploadFile = File(...),
    agent_id: Optional[str] = Form(None),
    broker_id: Optional[str] = Form(None),
    payment_mode: str = Form(...),
    amount_paid: Optional[float] = Form(None),
    payment_reference: Optional[str] = Form(None),
    user: User = Depends(current_user),
):
    if payment_mode not in PAYMENT_MODES:
        raise HTTPException(status_code=400, detail=f"Payment mode must be one of: {', '.join(PAYMENT_MODES)}")
    partial = payment_mode == "Partial Payment"
    if partial and (amount_paid is None or amount_paid <= 0):
        raise HTTPException(status_code=400, detail="Enter the amount paid for a partial payment")
    ext = Path(file.filename or "").suffix.lower()
    mime = file.content_type if file.content_type in ALLOWED else EXT_TO_MIME.get(ext)
    if mime not in ALLOWED:
        raise HTTPException(status_code=415, detail="Only PDF, JPG, PNG or WEBP files are supported")
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="File is empty")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="File exceeds 15 MB")

    if user.role == "agent":  # agents always upload into their own MIS
        owner = {"id": user.id, "full_name": user.full_name, "username": user.username, "branch": user.branch}
    elif not agent_id:
        raise HTTPException(status_code=400, detail="Select the agent this policy belongs to")
    else:
        allowed = {a.id: a for a in await agent_options_for(user)}
        if agent_id not in allowed:
            raise HTTPException(status_code=403, detail="Agent is outside your assigned scope")
        a = allowed[agent_id]
        owner = {"id": a.id, "full_name": a.full_name, "username": a.username, "branch": a.branch}

    broker = await _resolve_broker(broker_id, owner["id"])
    stored_data, stored_mime = await asyncio.to_thread(compress_document, data, mime)
    policy = Policy(
        file_name=file.filename or f"document{ALLOWED[mime]}",
        file_mime=stored_mime,
        file_size=len(stored_data),
        original_size=len(data),
        has_file=True,
        agent_id=owner["id"],
        broker_id=broker["id"] if broker else None,
        broker_name=broker["name"] if broker else None,
        agent_name=owner["full_name"],
        agent_code=owner["username"].upper(),
        branch=owner["branch"],
        uploaded_by_id=user.id,
        uploaded_by_name=user.full_name,
        history=[HistoryEntry(**history(user, "Uploaded", f"{file.filename} · {payment_mode}"))],
        payment_mode=payment_mode,
        payments=(
            [PaymentEntry(amount=round(amount_paid, 2), reference=payment_reference or None, paid_on=today_iso(), recorded_by_name=user.full_name)]
            if partial and amount_paid
            else []
        ),
    )
    try:
        stored = await put_object(object_path(owner["id"], policy.id, ALLOWED[stored_mime]), stored_data, stored_mime)
    except Exception:
        logger.exception("Object storage upload failed")
        raise HTTPException(status_code=502, detail="Could not store the document. Please retry.")
    await db.policies.insert_one(
        {**policy.to_doc(), "storage_path": stored["path"], "is_deleted": False}
    )
    _spawn_ocr(policy.id, stored["path"], stored_mime, original=data, original_mime=mime)
    return policy


@router.get("/policies", response_model=List[Policy])
async def list_policies(
    filters: PolicyFilters = Depends(), outstanding_only: bool = False, user: User = Depends(current_user)
):
    docs = await db.policies.find(filters.query(user), {"_id": 0, "history": 0}).sort("created_at", -1).to_list(1000)
    comm = await commissions_for(docs) if user.role == "admin" else {}
    policies = [Policy(**d, commission=comm.get(d["id"])) for d in docs]
    return [p for p in policies if (p.outstanding or 0) > 0] if outstanding_only else policies


@router.get("/policies/{id}", response_model=Policy)
async def get_policy(id: str, user: User = Depends(current_user)):
    return Policy(**await load_scoped(id, user))


@router.get("/policies/{id}/file")
async def get_policy_file(id: str, download: bool = False, user: User = Depends(current_user)):
    doc = await load_scoped(id, user)
    path = doc.get("storage_path")
    if not path:
        raise HTTPException(status_code=404, detail="No source document for this policy")
    try:
        data, content_type = await get_object(path)
    except Exception:
        logger.exception("Object storage download failed")
        raise HTTPException(status_code=502, detail="Could not load the document from storage")
    mime = doc.get("file_mime") or content_type
    stem = Path(doc.get("file_name") or "policy").stem
    label = (doc.get("fields") or {}).get("policy_number") or stem
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", str(label)) + ALLOWED.get(mime, Path(doc.get("file_name") or "").suffix)
    return Response(
        content=data,
        media_type=mime,
        headers={"Content-Disposition": f'{"attachment" if download else "inline"}; filename="{safe}"', "Cache-Control": "private, max-age=300"},
    )


@router.patch("/policies/{id}", response_model=Policy)
async def update_policy(id: str, body: PolicyUpdate, user: User = Depends(current_user)):
    doc = await load_scoped(id, user)
    if doc["status"] == "processing":
        raise HTTPException(status_code=409, detail="OCR is still running for this document")
    if user.role == "agent" and doc["status"] == "verified":
        raise HTTPException(status_code=403, detail="Verified policies are locked. Raise a flag to request a correction.")
    old = PolicyFields(**doc.get("fields", {})).model_dump()
    new = derive_fields(body.fields.model_dump())
    changed = [k for k in new if k not in DERIVED_FIELDS and new[k] != old.get(k)]
    dup = await find_duplicate(new, id)
    if dup and doc["status"] != "duplicate":
        raise HTTPException(status_code=409, detail=f"Duplicate — {dup[1]}")
    status = "extracted" if doc["status"] in ("failed", "duplicate") and not dup else doc["status"]
    note = (body.note + " · " if body.note else "") + (f"Changed: {', '.join(changed)}" if changed else "No field changes")
    dup_set = {"duplicate_reason": dup[1], "duplicate_of_id": dup[0]} if dup else {"duplicate_reason": None, "duplicate_of_id": None}
    await db.policies.update_one(
        {"id": id},
        {"$set": {"fields": new, "status": status, "updated_at": now_utc(), **dup_set}, "$push": {"history": history(user, "Corrected", note)}},
    )
    return Policy(**await load_scoped(id, user))


@router.post("/policies/{id}/verify", response_model=Policy)
async def verify_policy(id: str, body: VerifyIn, user: User = Depends(reviewers)):
    doc = await load_scoped(id, user)
    if doc["status"] == "processing":
        raise HTTPException(status_code=409, detail="OCR is still running for this document")
    f = doc.get("fields", {})
    if not f.get("policy_number") or not f.get("gross_premium"):
        raise HTTPException(status_code=400, detail="Policy number and gross premium are required before verification")
    dup = await find_duplicate(f, id)
    if dup or doc["status"] == "duplicate":
        raise HTTPException(status_code=409, detail=f"Duplicate — {dup[1] if dup else doc.get('duplicate_reason')}")
    now = now_utc()
    await db.policies.update_one(
        {"id": id},
        {
            "$set": {"status": "verified", "verified_by_name": user.full_name, "verified_at": now, "flag_reason": None, "updated_at": now},
            "$push": {"history": history(user, "Verified", body.note)},
        },
    )
    return Policy(**await load_scoped(id, user))


@router.post("/policies/{id}/flag", response_model=Policy)
async def flag_policy(id: str, body: FlagIn, user: User = Depends(current_user)):
    doc = await load_scoped(id, user)
    if doc["status"] == "processing":
        raise HTTPException(status_code=409, detail="OCR is still running for this document")
    await db.policies.update_one(
        {"id": id},
        {"$set": {"status": "flagged", "flag_reason": body.reason, "updated_at": now_utc()}, "$push": {"history": history(user, "Flagged", body.reason)}},
    )
    return Policy(**await load_scoped(id, user))


@router.post("/policies/{id}/reextract", response_model=Policy)
async def reextract_policy(id: str, user: User = Depends(current_user)):
    doc = await load_scoped(id, user)
    if user.role == "agent" and doc["status"] == "verified":
        raise HTTPException(status_code=403, detail="Verified policies are locked")
    path = doc.get("storage_path")
    if not path:
        raise HTTPException(status_code=400, detail="No source document to re-extract")
    await db.policies.update_one(
        {"id": id},
        {"$set": {"status": "processing", "ocr_error": None, "updated_at": now_utc()}, "$push": {"history": history(user, "Re-extraction requested")}},
    )
    _spawn_ocr(id, path, doc["file_mime"])
    return Policy(**await load_scoped(id, user))


@router.delete("/policies/{id}")
async def delete_policy(id: str, user: User = Depends(require_roles("admin"))):
    await load_scoped(id, user)
    # storage has no delete API — soft-delete the record; the object stays orphaned but unreachable
    await db.policies.update_one(
        {"id": id}, {"$set": {"is_deleted": True, "updated_at": now_utc()}, "$push": {"history": history(user, "Deleted")}}
    )
    return {"ok": True}


@router.post("/policies/reassign-agent", response_model=ReassignResult)
async def reassign_agent(body: ReassignAgentIn, user: User = Depends(reviewers)):
    """Correct the agent a policy belongs to (bulk). MIS + outstanding follow the new agent automatically."""
    allowed = {a.id: a for a in await agent_options_for(user)}
    a = allowed.get(body.agent_id)
    if not a:
        raise HTTPException(status_code=403, detail="Agent is outside your assigned scope")
    code = a.username.upper()
    updated = 0
    for pid in dict.fromkeys(body.policy_ids):
        doc = await load_scoped(pid, user)
        if doc.get("agent_id") == a.id:
            continue
        note = f"{doc.get('agent_name')} ({doc.get('agent_code')}) → {a.full_name} ({code})"
        await db.policies.update_one(
            {"id": pid},
            {"$set": {"agent_id": a.id, "agent_name": a.full_name, "agent_code": code, "branch": a.branch, "updated_at": now_utc()},
             "$push": {"history": history(user, "Agent changed", note)}},
        )
        updated += 1
    return ReassignResult(updated=updated, agent_name=a.full_name, agent_code=code)


@router.post("/policies/set-broker", response_model=ReassignResult)
async def bulk_set_broker(body: BulkBrokerIn, user: User = Depends(reviewers)):
    b = await db.brokers.find_one({"id": body.broker_id}, {"_id": 0, "id": 1, "name": 1}) if body.broker_id else None
    if body.broker_id and not b:
        raise HTTPException(status_code=404, detail="Broker not found")
    updated = 0
    for pid in dict.fromkeys(body.policy_ids):
        doc = await load_scoped(pid, user)
        if doc.get("broker_id") == (b["id"] if b else None):
            continue
        await db.policies.update_one(
            {"id": pid},
            {"$set": {"broker_id": b["id"] if b else None, "broker_name": b["name"] if b else None, "updated_at": now_utc()},
             "$push": {"history": history(user, "Broker changed", f"{doc.get('broker_name') or 'None'} → {b['name'] if b else 'None'}")}},
        )
        updated += 1
    return ReassignResult(updated=updated, agent_name=b["name"] if b else "No broker", agent_code="")


@router.put("/policies/{id}/broker", response_model=Policy)
async def set_broker(id: str, body: BrokerSetIn, user: User = Depends(reviewers)):
    doc = await load_scoped(id, user)
    b = await db.brokers.find_one({"id": body.broker_id}, {"_id": 0, "id": 1, "name": 1}) if body.broker_id else None
    if body.broker_id and not b:
        raise HTTPException(status_code=404, detail="Broker not found")
    await db.policies.update_one(
        {"id": id},
        {"$set": {"broker_id": b["id"] if b else None, "broker_name": b["name"] if b else None, "updated_at": now_utc()},
         "$push": {"history": history(user, "Broker changed", f"{doc.get('broker_name') or 'None'} → {b['name'] if b else 'None'}")}},
    )
    return Policy(**await load_scoped(id, user))


@router.put("/policies/{id}/payment-mode", response_model=Policy)
async def set_payment_mode(id: str, body: PaymentModeIn, user: User = Depends(reviewers)):
    doc = await load_scoped(id, user)
    if doc.get("payment_mode") == body.payment_mode:
        return Policy(**doc)
    await db.policies.update_one(
        {"id": id},
        {"$set": {"payment_mode": body.payment_mode, "updated_at": now_utc()},
         "$push": {"history": history(user, "Payment mode changed", f"{doc.get('payment_mode') or 'Not set'} → {body.payment_mode}")}},
    )
    return Policy(**await load_scoped(id, user))


@router.post("/policies/{id}/payments", response_model=Policy)
async def record_payment(id: str, body: PaymentIn, user: User = Depends(reviewers)):
    policy = Policy(**await load_scoped(id, user))
    if policy.payment_mode not in OUTSTANDING_MODES:
        raise HTTPException(status_code=400, detail="Payments can only be recorded on Cash or Partial Payment policies")
    if policy.outstanding is None:
        raise HTTPException(status_code=400, detail="Total premium is not captured yet — save it before recording payments")
    if body.amount > policy.outstanding + 0.5:
        raise HTTPException(status_code=400, detail=f"Amount exceeds the outstanding balance of {policy.outstanding:,.2f}")
    entry = PaymentEntry(amount=round(body.amount, 2), reference=body.reference or None, paid_on=body.paid_on or today_iso(), recorded_by_name=user.full_name)
    await db.policies.update_one(
        {"id": id},
        {"$push": {"payments": entry.model_dump(), "history": history(user, "Payment recorded", f"₹{entry.amount:,.2f}" + (f" · {entry.reference}" if entry.reference else ""))},
         "$set": {"updated_at": now_utc()}},
    )
    return Policy(**await load_scoped(id, user))


@router.delete("/policies/{id}/payments/{entry_id}", response_model=Policy)
async def delete_payment(id: str, entry_id: str, user: User = Depends(require_roles("admin"))):
    await load_scoped(id, user)
    res = await db.policies.update_one(
        {"id": id, "payments.id": entry_id},
        {"$pull": {"payments": {"id": entry_id}}, "$push": {"history": history(user, "Payment removed")}, "$set": {"updated_at": now_utc()}},
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Payment entry not found")
    return Policy(**await load_scoped(id, user))


@router.get("/outstanding", response_model=List[AgentOutstanding])
async def agent_outstanding(filters: PolicyFilters = Depends(), user: User = Depends(current_user)):
    query = filters.query(user)
    modes = {"payment_mode": {"$in": OUTSTANDING_MODES}}
    query = {"$and": [query, modes]} if query else modes
    docs = await db.policies.find(query, {"_id": 0, "history": 0}).to_list(20000)
    agg: dict = {}
    for p in (Policy(**d) for d in docs):
        a = agg.setdefault(p.agent_id, AgentOutstanding(agent_id=p.agent_id, agent_name=p.agent_name, agent_code=p.agent_code, branch=p.branch, policies=0, total_premium=0, paid=0, outstanding=0))
        a.policies += 1
        a.total_premium = round(a.total_premium + (p.fields.gross_premium or 0), 2)
        a.paid = round(a.paid + (p.amount_paid or 0), 2)
        a.outstanding = round(a.outstanding + (p.outstanding or 0), 2)
    return sorted(agg.values(), key=lambda a: a.outstanding, reverse=True)


@router.get("/stats", response_model=Stats)
async def stats(filters: PolicyFilters = Depends(), user: User = Depends(current_user)):
    docs = await db.policies.find(filters.query(user), {"_id": 0, "status": 1, "fields": 1, "agent_name": 1, "payment_mode": 1, "payments": 1}).to_list(20000)
    counts: dict = defaultdict(int)
    groups: dict = {k: defaultdict(lambda: [0, 0.0]) for k in ("lob", "insurer", "agent", "age")}
    gross = net = owed = 0.0
    for d in docs:
        counts[d["status"]] += 1
        f = d.get("fields") or {}
        g = float(f.get("gross_premium") or 0)
        owed += payment_totals(d.get("payment_mode"), d.get("payments") or [], f.get("gross_premium"))[1] or 0
        gross += g
        net += float(f.get("net_premium") or 0)
        for key, label in (("lob", f.get("lob")), ("insurer", f.get("insurer")), ("agent", d.get("agent_name"))):
            bucket = groups[key][label or "Unclassified"]
            bucket[0] += 1
            bucket[1] += g
        if f.get("vehicle_age_bucket"):
            groups["age"][f["vehicle_age_bucket"]][0] += 1
            groups["age"][f["vehicle_age_bucket"]][1] += g

    def top(key: str, n: int) -> List[GroupStat]:
        items = sorted(groups[key].items(), key=lambda kv: kv[1][1], reverse=True)[:n]
        return [GroupStat(label=k, count=v[0], premium=round(v[1], 2)) for k, v in items]

    return Stats(
        total=len(docs),
        processing=counts["processing"],
        extracted=counts["extracted"],
        flagged=counts["flagged"],
        verified=counts["verified"],
        failed=counts["failed"],
        gross_premium=round(gross, 2),
        net_premium=round(net, 2),
        by_lob=top("lob", 8),
        by_insurer=top("insurer", 6),
        by_agent=top("agent", 6),
        outstanding=round(owed, 2),
        by_age=[GroupStat(label=b, count=groups["age"][b][0], premium=round(groups["age"][b][1], 2)) for b in AGE_BUCKETS],
    )


# ---------- storage & OCR usage (admin) ----------
OCR_COST_PER_POLICY_INR = float(os.environ.get("OCR_COST_PER_POLICY_INR") or 0.85)  # est. Gemini Flash tokens per document


def _ocr_events(doc: dict) -> list:
    return [h for h in doc.get("history") or [] if h.get("by_role") == "system" and h.get("action") in ("OCR extracted", "OCR failed")]


@router.get("/usage", response_model=Usage)
async def usage(user: User = Depends(require_roles("admin"))):
    docs = await db.policies.find(
        {"seeded": {"$ne": True}, "has_file": True},
        {"_id": 0, "branch": 1, "agent_name": 1, "agent_code": 1, "file_size": 1, "original_size": 1, "created_at": 1, "history": 1},
    ).to_list(100000)
    month_key = today_iso()[:7]
    groups: dict = {"branch": {}, "agent": {}}
    months: dict = defaultdict(lambda: [0, 0])  # ocr_runs, files
    tot = {"stored": 0, "orig": 0, "small": 0, "runs": 0, "ok": 0, "fail": 0, "month": 0}
    for d in docs:
        stored = int(d.get("file_size") or 0)
        orig = int(d.get("original_size") or stored)
        events = _ocr_events(d)
        tot["stored"] += stored
        tot["orig"] += orig
        tot["small"] += 1 if 0 < stored <= TARGET_BYTES else 0
        tot["ok"] += sum(1 for e in events if e["action"] == "OCR extracted")
        tot["runs"] += len(events)
        if isinstance(d.get("created_at"), datetime):
            months[d["created_at"].strftime("%Y-%m")][1] += 1
        for e in events:
            m = e["at"].strftime("%Y-%m") if isinstance(e.get("at"), datetime) else month_key
            months[m][0] += 1
            tot["month"] += 1 if m == month_key else 0
        for key, label in (("branch", d.get("branch") or "Unassigned"), ("agent", f'{d.get("agent_name")} ({d.get("agent_code")})')):
            g = groups[key].setdefault(label, [0, 0, 0, 0])
            g[0] += 1
            g[1] += stored
            g[2] += orig
            g[3] += len(events)
    tot["fail"] = tot["runs"] - tot["ok"]
    cost = lambda runs: round(runs * OCR_COST_PER_POLICY_INR, 2)  # noqa: E731

    def group_list(key: str) -> List[UsageGroup]:
        items = sorted(groups[key].items(), key=lambda kv: kv[1][0], reverse=True)
        return [UsageGroup(label=k, files=v[0], stored_bytes=v[1], original_bytes=v[2], ocr_runs=v[3], est_cost_inr=cost(v[3])) for k, v in items]

    return Usage(
        total_files=len(docs),
        stored_bytes=tot["stored"],
        original_bytes=tot["orig"],
        avg_file_bytes=tot["stored"] // len(docs) if docs else 0,
        under_target_files=tot["small"],
        target_bytes=TARGET_BYTES,
        ocr_runs=tot["runs"],
        ocr_success=tot["ok"],
        ocr_failed=tot["fail"],
        ocr_runs_this_month=tot["month"],
        est_cost_per_policy_inr=OCR_COST_PER_POLICY_INR,
        est_cost_total_inr=cost(tot["runs"]),
        est_cost_this_month_inr=cost(tot["month"]),
        model_name=MODEL_NAME,
        by_month=[UsageMonth(month=m, ocr_runs=v[0], files=v[1], est_cost_inr=cost(v[0])) for m, v in sorted(months.items())[-12:]],
        by_branch=group_list("branch"),
        by_agent=group_list("agent"),
    )


# ---------- MIS export ----------
# Exact column order / headings from the business "MIS order.xlsx" (A..AH), then payment balance columns
MIS_COLUMNS = [
    ("Agent Name", "agent_name"), ("Broker /Agent/POS Name", "f.intermediary_name"), ("Insurer", "f.insurer"), ("LOB", "f.lob"),
    ("Risk Start Date", "f.start_date"), ("Customer Name", "f.insured_name"), ("Policy No", "f.policy_number"),
    ("Registration No", "f.vehicle_reg_no"), ("Year of Mfg", "f.mfg_year"), ("Product", "f.product_name"), ("Make", "f.make"),
    ("Model", "f.model"), ("IDV /Sum Insured", "x.idv_si"), ("OD Premium (incl. Add-ons)", "f.od_premium"),
    ("TP Premium (Liability)", "f.tp_premium"), ("Net Premium", "f.net_premium"), ("GST", "f.gst"), ("Total Premium", "f.gross_premium"),
    ("Risk Start Date", "f.start_date"), ("GVW (kg)", "f.gvw"), ("CC", "f.cc"), ("Seating Capacity", "f.seating_capacity"),
    ("Year of Mfg", "f.mfg_year"), ("RTO", "f.rto"), ("Vehicle Age (Yrs)", "f.vehicle_age_years"), ("Vehicle Age Band", "f.vehicle_age_bucket"),
    ("Status", "status"), ("Flag Reason", "flag_reason"), ("Verified By", "verified_by_name"), ("Verified At", "verified_at"),
    ("Uploaded By", "uploaded_by_name"), ("Uploaded At", "created_at"), ("Source File", "file_name"),
    ("Agent Premium Paid By Cash /Online/Cheque/Partial Payment", "payment_mode"),
    ("Amount Paid", "x.amount_paid"), ("Agent Outstanding Premium", "x.outstanding"),
]
COMMISSION_HEADERS = ["Broker (Commission)", "Product Category", "Commission Receivable", "Commission Payable", "Commission Margin"]
TEXT_COLS = {"Policy No", "Registration No", "Year of Mfg", "RTO"}  # long numeric ids must stay text in Excel
MONEY_COLS = {"Commission Receivable", "Commission Payable", "Commission Margin", "IDV /Sum Insured", "OD Premium (incl. Add-ons)", "TP Premium (Liability)", "Net Premium", "GST", "Total Premium", "Amount Paid", "Agent Outstanding Premium"}



def _extras(doc: dict, fields: dict) -> dict:
    paid, owed = payment_totals(doc.get("payment_mode"), doc.get("payments") or [], fields.get("gross_premium"))
    idv_si = fields.get("idv") if fields.get("lob") == "Motor" else fields.get("sum_insured")
    return {"idv_si": idv_si if idv_si is not None else fields.get("sum_insured"), "amount_paid": paid, "outstanding": owed}


def _cell(doc: dict, fields: dict, key: str, extras: dict):
    if key.startswith("x."):
        return extras.get(key[2:])
    if key == "f.gvw" and fields.get("vehicle_class") in GVW_NA_CLASSES:
        return "NA"
    if key.startswith("f."):
        return fields.get(key[2:])
    v = doc.get(key)
    if isinstance(v, datetime):
        return v.replace(tzinfo=timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    return v


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1E293B")


def _style_header(ws) -> None:
    for c in ws[1]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(vertical="center")


def _money_format(ws, col: str) -> None:
    for cell in ws[col][1:]:
        cell.number_format = "#,##0.00"


def _csv_bytes(headers: list, rows: list) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    text_idx = {i for i, h in enumerate(headers) if h in TEXT_COLS}

    def csv_value(i: int, v):
        if v is None:
            return ""
        if i in text_idx and str(v).isdigit() and len(str(v)) > 11:
            return f'="{v}"'  # stops Excel turning long policy numbers into 5.005E+17
        return v

    w.writerows([[csv_value(i, v) for i, v in enumerate(r)] for r in rows])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def _age_summary(derived: list) -> dict:
    agg = {b: [0, 0.0, 0.0, 0.0] for b in AGE_BUCKETS}
    for f in derived:
        bucket = f.get("vehicle_age_bucket")
        if bucket:
            a = agg[bucket]
            a[0] += 1
            a[1] += f.get("gross_premium") or 0
            a[2] += f.get("od_premium") or 0
            a[3] += f.get("tp_premium") or 0
    return agg


def _write_mis_sheet(ws, headers: list, rows: list) -> None:
    ws.title = "Policy MIS"
    ws.append(headers)
    _style_header(ws)
    for r in rows:
        ws.append(r)
    for idx, h in enumerate(headers, start=1):
        col = ws.cell(row=1, column=idx).column_letter
        longest = max([len(h)] + [len(str(c.value)) for c in ws[col][1:] if c.value is not None])
        ws.column_dimensions[col].width = max(12, min(48, longest + 3))
        if h in MONEY_COLS:
            _money_format(ws, col)
        if h in TEXT_COLS:
            for cell in ws[col][1:]:
                cell.number_format = "@"
                if cell.value is not None:
                    cell.value = str(cell.value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_age_sheet(ws, derived: list) -> None:
    ws.append(["Vehicle Age Band", "Policies", "Total Premium", "OD Premium", "TP Premium"])
    _style_header(ws)
    for bucket, values in _age_summary(derived).items():
        ws.append([bucket, *values])
    for col, width in zip("ABCDE", (22, 12, 18, 16, 16)):
        ws.column_dimensions[col].width = width
        if col in "CDE":
            _money_format(ws, col)


def _xlsx_bytes(headers: list, rows: list, derived: list) -> io.BytesIO:
    wb = Workbook()
    _write_mis_sheet(wb.active, headers, rows)
    _write_age_sheet(wb.create_sheet("Vehicle Age Summary"), derived)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _export_filename(user: User, filters: PolicyFilters, fmt: str) -> str:
    scope = user.username.upper() if user.role != "admin" else ("AGENT" if filters.agent_id else "ALL")
    return f"MIS_{scope}_{today_iso()}.{fmt}"


@router.get("/mis/export")
async def export_mis(
    format: Literal["csv", "xlsx"] = Query("xlsx"),
    filters: PolicyFilters = Depends(),
    outstanding_only: bool = False,
    user: User = Depends(current_user),
):
    docs = await db.policies.find(filters.query(user), {"_id": 0, "history": 0}).sort("created_at", -1).to_list(50000)
    if outstanding_only:
        docs = [d for d in docs if (payment_totals(d.get("payment_mode"), d.get("payments") or [], (d.get("fields") or {}).get("gross_premium"))[1] or 0) > 0]
    derived = [derive_fields(d.get("fields") or {}) for d in docs]
    rows = [[_cell(d, f, key, _extras(d, f)) for _, key in MIS_COLUMNS] for d, f in zip(docs, derived)]
    headers = [h for h, _ in MIS_COLUMNS]
    if user.role == "admin":  # commission columns appended after the standard MIS — admin only
        comm = await commissions_for(docs)
        headers += COMMISSION_HEADERS
        for r, d, f in zip(rows, docs, derived):
            c = comm.get(d["id"])
            r += [d.get("broker_name"), f.get("product_category"), c and c.receivable, c and c.payable, c and c.margin]
    disposition = {"Content-Disposition": f'attachment; filename="{_export_filename(user, filters, format)}"'}
    if format == "csv":
        return StreamingResponse(iter([_csv_bytes(headers, rows)]), media_type="text/csv", headers=disposition)
    return StreamingResponse(
        _xlsx_bytes(headers, rows, derived),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=disposition,
    )
